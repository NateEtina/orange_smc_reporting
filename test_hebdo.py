"""End-to-end tests for the weekly report engine (Feature 1: P1/P2 MTTR).

Run:  python test_hebdo.py
"""

from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta

from hebdo import (
    HebdoError,
    _domain_row,
    analyze,
    format_mttr,
    generate,
    parse_day,
    parse_duration,
    parse_week,
    render_excel,
    render_pptx,
    resolve_columns,
    resolve_week,
    week_bounds,
    week_value_matches,
)

FAILURES: list[str] = []
CHECKS = 0


def check(condition, message: str) -> None:
    global CHECKS
    CHECKS += 1
    if not condition:
        FAILURES.append(message)


def load_xlsx(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


def load_pptx(data: bytes):
    from pptx import Presentation

    return Presentation(io.BytesIO(data))


def cell_fill(cell) -> str:
    fill = cell.fill
    if fill is None or not fill.patternType:
        return ""
    return str(fill.fgColor.rgb or "")


# --------------------------------------------------------------------------
# Fixtures (semicolon-delimited, DMY dates in brise, ISO dates in GNOC;
# the 'Duree' column carries the resolution time in hours)
# --------------------------------------------------------------------------

BRISE_CSV = "\n".join([
    "Ticket;Date creation;Date resolution;semaine creation;identifiant EDS active;"
    "priorite traitement;SLA;Duree",
    "T1;21/09/2026 08:00;21/09/2026 10:00;S39;505010;P1;GOOD;2",
    "T2;21/09/2026 09:00;21/09/2026 14:30;S39;555402;P1;BAD;5,5",
    "T3;22/09/2026 08:00;;S39;610136;P2;En Cours;",
    "T4;22/09/2026 08:00;22/09/2026 09:30;S39;610136;P2;GOOD;1:30",
    "T5;23/09/2026 07:00;23/09/2026 09:00;39;610136;P3;GOOD;2",
    "T6;24/09/2026 08:00;24/09/2026 09:00;S39;999999;P1;GOOD;1",
    "T7;18/09/2026 08:00;18/09/2026 09:00;S38;505010;P1;GOOD;1",
    "T8;25/09/2026 08:00;25/09/2026 09:00;S39;505010;P5;GOOD;1",
]) + "\n"

GNOC_CSV = "\n".join([
    "Ticket;semaine creation;priorite traitement;valeur champ 2;Statut SLA Retablissement;"
    "date creation;date resolution;Duree",
    "G1;S39;P1;core CS;GOOD;2026-09-21 08:00;2026-09-21 09:00;1",
    "G2;S39;P1;core CS;;2026-09-22 10:00;;",
    "G3;S39;P2;Core IP;GOOD;2026-09-22 09:00;2026-09-22 12:00;3",
    "G4;S39;P2;IN;GOOD;2026-09-22 08:00;2026-09-22 10:00;",
    "G5;S39;P2;VAS;BAD;2026-09-23 08:00;2026-09-24 14:00;30:00",
    "G6;S39;P3;Core IP;BAD;2026-09-24 09:00;2026-09-24 10:00;1",
    "G7;S39;P2;Storage;GOOD;2026-09-24 09:00;2026-09-24 10:00;1",
    "G8;S38;P2;core CS;GOOD;2026-09-18 08:00;2026-09-18 09:00;1",
    "G9;S39;P4;core CS;Non évalué;2026-09-25 08:00;2026-09-25 09:00;1",
    "G10;S39;P1;Core IP;GOOD;2026-09-24 10:00;2026-09-24 08:00;-2",
    "G11;S39;P1;TRANS;GOOD;2026-09-23 08:00;2026-09-23 12:00;abc",
    "G12;S39;P2;IN/TRANS/VAS;GOOD;2026-09-22 08:00;2026-09-22 08:00;0",
    "G13;S39;P4;Core IP;;2026-09-24 08:00;;",
]) + "\n"

BRISE = ("brise_smc_week39.csv", BRISE_CSV.encode("utf-8"))
GNOC = ("GNOC_week39.csv", GNOC_CSV.encode("utf-8"))


def run_tests() -> None:
    # ================= week management ====================================
    check(parse_week("S39") == 39, "parse_week S39")
    check(parse_week("s 39") == 39, "parse_week 's 39'")
    check(parse_week(39) == 39, "parse_week int")
    check(parse_week("39") == 39, "parse_week '39'")
    check(parse_week("Semaine 39") == 39, "parse_week 'Semaine 39'")
    for bad in ("S54", "abc", "", "S0"):
        try:
            parse_week(bad)
            check(False, f"parse_week should reject {bad!r}")
        except HebdoError:
            check(True, "")

    check(parse_day("2026-09-21") == date(2026, 9, 21), "parse_day ISO")
    check(parse_day("21/09/2026") == date(2026, 9, 21), "parse_day DMY")
    try:
        parse_day("not-a-date")
        check(False, "parse_day should reject junk")
    except HebdoError:
        check(True, "")

    check(resolve_week("S39") == (39, date.today().isocalendar().year), "resolve_week explicit")
    check(resolve_week(None, "2026-09-21") == (39, 2026), "resolve_week from Monday 2026-09-21")
    check(resolve_week(None, "2026-09-27") == (39, 2026), "resolve_week from Sunday 2026-09-27")
    check(resolve_week(None, "2026-09-28") == (40, 2026), "resolve_week rolls to next ISO week")
    check(resolve_week() == tuple(date.today().isocalendar()[:2][::-1])
          or resolve_week() == (date.today().isocalendar()[1], date.today().isocalendar()[0]),
          "resolve_week defaults to current ISO week")
    check(week_bounds(39, 2026) == (date(2026, 9, 21), date(2026, 9, 27)),
          "week_bounds S39 2026 = Sep 21-27")

    # ---- resolution time ('Duree' column) ------------------------------
    check(parse_duration(2.5) == 2.5, "Duree numeric hours")
    check(parse_duration("36.5") == 36.5, "Duree decimal string")
    check(parse_duration("36,5") == 36.5, "Duree comma decimal string")
    check(parse_duration("2:30") == 2.5, "Duree HH:MM")
    check(parse_duration("36:30") == 36.5, "Duree HH:MM beyond 24 h")
    check(parse_duration("1:30:00") == 1.5, "Duree HH:MM:SS")
    check(parse_duration("4h") == 4.0, "Duree with trailing h")
    check(parse_duration("4.5 h") == 4.5, "Duree with ' h' suffix")
    check(parse_duration("5h30") == 5.5, "Duree compact '5h30'")
    check(parse_duration("5H30") == 5.5, "Duree compact '5H30'")
    check(parse_duration("5 h 30") == 5.5, "Duree spaced '5 h 30'")
    check(parse_duration("12h45m") == 12.75, "Duree compact '12h45m'")
    check(parse_duration("3:26:0") == 3 + 26 / 60, "Duree '3:26:0' single-digit seconds")
    check(parse_duration("0:6:59") == 6 / 60 + 59 / 3600, "Duree '0:6:59' single-digit minutes")
    check(parse_duration("1:18:0") == 1 + 18 / 60, "Duree '1:18:0' real export value")
    check(parse_duration("0:0:59") == 59 / 3600, "Duree '0:0:59'")
    check(parse_duration("0:60:0") is None, "Duree invalid minutes 60 -> None")
    check(parse_duration("-1110976:59:59") is None, "Duree garbage negative -> None")
    check(parse_duration(timedelta(hours=3, minutes=30)) == 3.5, "Duree Excel duration cell")
    check(parse_duration(time(6, 30)) == 6.5, "Duree Excel time cell")
    check(parse_duration(0) == 0.0, "Duree zero parses to 0")
    check(parse_duration(-2) == -2.0, "Duree negative parses")
    for blank in (None, "", "   ", "abc", "1:99", True):
        check(parse_duration(blank) is None, f"Duree blank/unparsable -> None: {blank!r}")

    # MTTR display: total hours as HH:MM, beyond 24 h included
    check(format_mttr(3.75) == "03:45", "format_mttr 3.75 -> 03:45")
    check(format_mttr(36.5) == "36:30", "format_mttr 36.5 -> 36:30 (>24 h)")
    check(format_mttr(4.1) == "04:06", "format_mttr 4.1 -> 04:06")
    check(format_mttr(100.75) == "100:45", "format_mttr 100.75 -> 100:45")

    # ---- composite domain mapping (IN / VAS / TRANS -> IN/TRANS/VAS) ----
    for value in ("IN", "in", "VAS", "vas", "TRANS", "trans", "IN/TRANS/VAS",
                  "in/trans/vas", "IN, VAS, TRANS", "  trans  "):
        check(_domain_row(value) == "IN/TRANS/VAS", f"composite domain: {value!r}")
    check(_domain_row("Core CS") == "Core CS", "domain 'Core CS' passthrough")
    check(_domain_row("core ip") == "Core IP", "domain case-insensitive")
    check(_domain_row("INCIDENT") is None, "'INCIDENT' is not the IN sub-domain")
    check(_domain_row("Storage") is None, "unknown domain stays unmapped")
    check(_domain_row("") is None and _domain_row(None) is None, "empty domain")

    # ---- week value matching (the 'semaine creation' column) -------------
    for value in ("S39", "s39", 39, 39.0, "39", "39.0", "Semaine 39", "2026-S39",
                  datetime(2026, 9, 24, 10, 30), date(2026, 9, 21), date(2026, 9, 27),
                  "24/09/2026", "2026-09-24"):
        check(week_value_matches(value, 39, 2026) is True, f"week match: {value!r}")
    for value in (None, "", "  ", "S38", 38, "2025-S39", "abc", "2026-09-20",
                  date(2025, 9, 24), True):
        check(week_value_matches(value, 39, 2026) is False, f"week mismatch: {value!r}")

    # ================= column resolution ==================================
    headers = ["Ticket", "Date creation", "Date resolution", "semaine creation",
               "identifiant EDS active", "priorite traitement", "SLA", "Duree"]
    cols = resolve_columns(headers, None, ("week", "eds", "priority", "sla"))
    check(cols["week"] == "semaine creation", "auto week column")
    check(cols["eds"] == "identifiant EDS active", "auto eds column")
    check(cols["priority"] == "priorite traitement", "auto priority column")
    check(cols["sla"] == "SLA", "auto sla column")
    check(cols["duration"] == "Duree", "auto duration column")
    check(cols["domain"] is None, "no domain column in a brise-like file")

    # GNOC exports spell the verdict column 'Statut SLA Retablissement'
    cols_gnoc = resolve_columns(["Ticket", "semaine creation", "priorite traitement",
                                 "valeur champ 2", "Statut SLA Retablissement", "Duree"],
                                None, ("week", "priority", "sla", "domain", "duration"))
    check(cols_gnoc["sla"] == "Statut SLA Retablissement",
          "GNOC 'Statut SLA Retablissement' maps to the sla role")

    cols2 = resolve_columns(["prio_custom", "semaine creation"],
                            {"priority": "prio_custom"}, ("week", "priority"))
    check(cols2["priority"] == "prio_custom", "override wins")

    try:
        resolve_columns(headers, {"priority": "nope"}, ("week", "priority"))
        check(False, "unknown override should raise")
    except HebdoError:
        check(True, "")
    try:
        resolve_columns(["Ticket", "semaine creation"], None,
                        ("week", "priority", "sla", "duration"))
        check(False, "missing required column should raise")
    except HebdoError as exc:
        check("priority" in str(exc) and "duration" in str(exc),
              f"error lists missing roles: {exc}")

    # accent-insensitive matching (xlsx sources often carry accents)
    accented = ["Ticket", "Semaine création", "Priorité de traitement", "Valeur champ 2",
                "SLA", "Date de création", "Date de résolution", "Durée",
                "Identifiant EDS actif"]
    cols3 = resolve_columns(accented, None,
                            ("week", "eds", "priority", "sla", "domain", "duration"))
    check(cols3["week"] == "Semaine création", "accented week column")
    check(cols3["priority"] == "Priorité de traitement", "accented priority column")
    check(cols3["domain"] == "Valeur champ 2", "accented domain column")
    check(cols3["duration"] == "Durée", "accented duration column (Durée)")
    check(cols3["eds"] == "Identifiant EDS actif", "accented eds column (actif/active)")

    # a file without any Duree column is refused up front
    try:
        analyze(("b.csv", b"Ticket;semaine creation;identifiant EDS active;"
                          b"priorite traitement;SLA\nX1;S39;505010;P1;GOOD\n"),
                GNOC, week="S39")
        check(False, "missing Duree column should raise")
    except HebdoError as exc:
        check("duration" in str(exc), f"missing Duree names the role: {exc}")

    # ================= analysis ===========================================
    metrics = analyze(BRISE, GNOC, week="S39")
    rows = metrics["rows"]

    check(metrics["week"] == 39 and metrics["week_label"] == "S39", "metrics week label")
    check(metrics["week_start"] == "2026-09-21" and metrics["week_end"] == "2026-09-27",
          "metrics week bounds")
    check(metrics["row_order"] == ["Core CS", "Core IP", "Core PS", "IN/TRANS/VAS", "IT"],
          "template row order")

    it = rows["IT"]
    check(it["tt"] == 5, f"IT tt {it['tt']}")
    check(it["rates"]["P1"] == 0.5, f"IT P1 rate {it['rates']['P1']}")
    check(it["rates"]["P2"] == 0.5, f"IT P2 rate {it['rates']['P2']}")
    check(it["rates"]["P3"] == 1.0, f"IT P3 rate {it['rates']['P3']}")
    check(it["rates"]["P4"] is None, f"IT P4 rate {it['rates']['P4']}")
    check((it["in_sla"], it["hors_sla"], it["en_cours"]) == (3, 1, 1),
          f"IT SLA tallies {(it['in_sla'], it['hors_sla'], it['en_cours'])}")
    check(it["prio_total"] == {"P1": 2, "P2": 2, "P3": 1, "P4": 0}, f"IT prio {it['prio_total']}")
    # MTTR = average Duree over resolved tickets (status rule), shown as HH:MM
    check(it["mttr"]["P1"] == "03:45" and it["mttr_ok"]["P1"] is True,
          f"IT MTTR P1 {it['mttr']['P1']} ok={it['mttr_ok']['P1']}")
    check(it["mttr"]["P2"] == "01:30" and it["mttr_ok"]["P2"] is True,
          f"IT MTTR P2 {it['mttr']['P2']} (Duree '1:30' read as 1.5 h)")
    check(it["ouvert"] == {"P1": 0, "P2": 1},
          f"IT ouvert {it['ouvert']} (brise 'En Cours' = open)")

    cs = rows["Core CS"]
    check(cs["tt"] == 3, f"Core CS tt {cs['tt']}")
    check(cs["rates"]["P1"] == 0.5, f"Core CS P1 rate {cs['rates']['P1']}")
    check(cs["rates"]["P4"] is None,
          "Core CS P4: no evaluated SLA verdict -> no data (not 0.00%)")
    check((cs["in_sla"], cs["hors_sla"], cs["en_cours"]) == (1, 0, 2),
          f"Core CS tallies {(cs['in_sla'], cs['hors_sla'], cs['en_cours'])}")
    check(cs["mttr"]["P1"] == "01:00" and cs["ouvert"]["P1"] == 1, "Core CS MTTR/ouvert P1")
    check(cs["mttr"]["P2"] is None, "Core CS has no P2 MTTR")

    ip = rows["Core IP"]
    check(ip["tt"] == 4, f"Core IP tt {ip['tt']}")
    check(ip["rates"]["P1"] == 1.0 and ip["rates"]["P2"] == 1.0,
          f"Core IP rates {ip['rates']}")
    check(ip["rates"]["P3"] == 0.0,
          "Core IP P3: evaluated BAD verdict -> real 0.00% (numeric)")
    check(ip["rates"]["P4"] is None,
          "Core IP P4: blank SLA everywhere -> empty (not 0.00%)")
    check(ip["mttr"]["P1"] is None and ip["mttr_ok"]["P1"] is None,
          "negative Duree: resolved ticket excluded from MTTR (Duree > 0) -> '-'")
    check(ip["ouvert"]["P1"] == 0,
          "resolved status wins: negative Duree no longer counts as open")
    check(ip["mttr"]["P2"] == "03:00" and ip["mttr_ok"]["P2"] is True,
          f"Core IP MTTR P2 {ip['mttr']['P2']}")

    vas = rows["IN/TRANS/VAS"]
    check(vas["tt"] == 4,
          f"IN/TRANS/VAS tt {vas['tt']} (IN + VAS + TRANS + literal composite)")
    check(abs(vas["rates"]["P2"] - 2 / 3) < 1e-9, f"IN/TRANS/VAS P2 rate {vas['rates']['P2']}")
    check(vas["mttr"]["P2"] == "30:00" and vas["mttr_ok"]["P2"] is False,
          f"IN/TRANS/VAS MTTR P2 {vas['mttr']['P2']} over 8h target - sum/count with the "
          "Duree > 0 filter: G5 30 h only (G4 blank, G12 0 h excluded -> 30/1, not 10)")
    check(vas["mttr"]["P1"] is None and vas["mttr_ok"]["P1"] is None,
          "unparsable Duree on a resolved ticket excluded -> no P1 MTTR ('-')")
    check(vas["ouvert"] == {"P1": 0, "P2": 0},
          f"vas ouvert {vas['ouvert']} (all resolved by status, none open)")

    ps = rows["Core PS"]
    check(ps["tt"] is None and all(v is None for v in ps["rates"].values()),
          "missing domain row stays empty (null counts, dashes)")
    check(ps["mttr"]["P1"] is None and ps["ouvert"] == {"P1": None, "P2": None},
          "missing domain: no MTTR, null ouvert -> '-' (no data)")

    # sources & warnings
    src_b, src_g = metrics["sources"]
    check(src_b["rows"] == 8 and src_b["matched_week"] == 7, f"brise counts {src_b}")
    check(src_b["excluded_eds"] == 1 and src_b["excluded_priority"] == 1 and src_b["used"] == 5,
          f"brise exclusions {src_b}")
    check(src_g["rows"] == 13 and src_g["matched_week"] == 12, f"GNOC counts {src_g}")
    check(src_g["excluded_domain"] == 1 and src_g["used"] == 11, f"GNOC exclusions {src_g}")

    warnings = "\n".join(metrics["warnings"])
    check("unknown priority: P5" in warnings, f"warning unknown priority: {metrics['warnings']}")
    check("active EDS unit is not in 505010" in warnings, "warning EDS exclusion")
    check("outside the known sub-domains: Storage" in warnings, "warning unknown domain")
    check("Non évalué" in warnings, "warning unknown SLA value")
    check("negative resolution time" in warnings, "warning negative duration")
    check("could not be parsed" in warnings, "warning unparsable duration")
    check("resolved ticket(s) without a resolution time" in warnings
          and "excluded from MTTR" in warnings,
          f"warning resolved-without-Duree excluded from MTTR: {metrics['warnings']}")
    check(metrics["columns"]["brise"]["duration"] == "Duree", "detected duration column")
    check(metrics["columns"]["gnoc"]["sla"] == "Statut SLA Retablissement",
          "detected GNOC status column under the sla role")
    check(metrics["columns"]["gnoc"]["domain"] == "valeur champ 2", "detected domain column")

    # default week = current ISO week (no explicit input)
    default_metrics = analyze(BRISE, GNOC)
    today = date.today().isocalendar()
    check(default_metrics["week"] == today[1] and default_metrics["year"] == today[0],
          "default week resolves to today's ISO week")

    # week with no data: analysis still works, warnings explain, label dynamic
    off = analyze(BRISE, GNOC, week="S40")
    check(off["week_label"] == "S40", "off-week label")
    check(any("no row matches the target week" in w for w in off["warnings"]),
          "off-week warning present")
    check(off["rows"]["IT"]["tt"] is None, "off-week IT empty (null tt -> '-')")

    # cp1252-encoded source (accents must survive)
    metrics_cp = analyze(BRISE, ("GNOC.csv", GNOC_CSV.encode("cp1252")), week="S39")
    check(metrics_cp["rows"]["Core CS"]["en_cours"] == 2, "cp1252 source processed")

    # ---- MTTR SLA thresholds: strictly above the target is BAD -----------
    mini_b = ("b.csv", ("Ticket;semaine creation;identifiant EDS active;priorite traitement;"
                        "SLA;Duree\n"
                        "X1;S39;505010;P1;GOOD;3\n"
                        "X2;S39;505010;P1;BAD;5\n").encode("utf-8"))
    mini_g = ("g.csv", ("Ticket;semaine creation;priorite traitement;valeur champ 2;SLA;Duree\n"
                        "Y1;S39;P1;core CS;GOOD;4\n"
                        "Y2;S39;P1;core CS;GOOD;4,0001\n").encode("utf-8"))
    boundary = analyze(mini_b, mini_g, week="S39")
    check(boundary["rows"]["IT"]["mttr"]["P1"] == "04:00"
          and boundary["rows"]["IT"]["mttr_ok"]["P1"] is True,
          "MTTR exactly 4.0 h is NOT bad")
    check(boundary["rows"]["Core CS"]["mttr_ok"]["P1"] is False
          and boundary["rows"]["Core CS"]["mttr"]["P1"] == "04:00",
          "MTTR above 4.0 h is BAD (P1)")

    # ---- Ouvert/MTTR population: file-specific status rules --------------
    # brise: open <=> status "En cours" (case-insensitive + trimmed)
    # GNOC : open <=> blank 'Statut SLA Retablissement'
    # MTTR subset = resolved AND Duree > 0; anything else is excluded.
    status_b = ("b.csv", ("Ticket;semaine creation;identifiant EDS active;priorite traitement;"
                          "SLA;Duree\n"
                          "X1;S39;505010;P1;En cours;9\n"
                          "X2;S39;505010;P1;  GOOD  ;4\n"
                          "X3;S39;505010;P2;;2\n"
                          "X4;S39;505010;P2; EN COURS ;7\n").encode("utf-8"))
    status_g = ("g.csv", ("Ticket;semaine creation;priorite traitement;valeur champ 2;"
                          "Statut SLA Retablissement;Duree\n"
                          "Y1;S39;P1;core CS;;7\n"
                          "Y2;S39;P1;core CS;GOOD;6\n"
                          "Y3;S39;P2;core IP;BAD;\n").encode("utf-8"))
    st_m = analyze(status_b, status_g, week="S39")
    st_it = st_m["rows"]["IT"]
    check(st_it["ouvert"] == {"P1": 1, "P2": 1},
          f"brise ouvert by status {st_it['ouvert']} (blank SLA = resolved)")
    check(st_it["mttr"]["P1"] == "04:00" and st_it["mttr_ok"]["P1"] is True,
          "brise MTTR ignores the open ticket's Duree (9 h not averaged)")
    check(st_it["mttr"]["P2"] == "02:00",
          "brise: blank SLA = resolved -> Duree enters MTTR")
    check(st_m["rows"]["Core CS"]["ouvert"]["P1"] == 1
          and st_m["rows"]["Core CS"]["mttr"]["P1"] == "06:00",
          "GNOC: blank status = open (7 h ignored), GOOD = resolved")
    check(st_m["rows"]["Core IP"]["mttr"]["P2"] is None
          and st_m["rows"]["Core IP"]["ouvert"]["P2"] == 0,
          "resolved GNOC ticket with blank Duree excluded from MTTR -> '-'")
    check(any("without a resolution time" in w for w in st_m["warnings"]),
          "resolved-without-Duree warning raised")

    # ---- final MTTR formula: SUM(Duree) / COUNT (resolved AND Duree > 0) --
    formula_b = ("b.csv", ("Ticket;semaine creation;identifiant EDS active;priorite traitement;"
                           "SLA;Duree\n"
                           "F1;S39;505010;P1;GOOD;3\n"
                           "F2;S39;505010;P1;GOOD;5\n"
                           "F3;S39;505010;P1;GOOD;0\n"
                           "F4;S39;505010;P1;GOOD;-4\n"
                           "F5;S39;505010;P1;En cours;99\n"
                           "F6;S39;505010;P2;GOOD;2\n"
                           "F7;S39;505010;P2;En cours;4\n").encode("utf-8"))
    formula_g = ("g.csv", ("Ticket;semaine creation;priorite traitement;valeur champ 2;"
                           "Statut SLA Retablissement;Duree\n"
                           "F8;S39;P2;core CS;;8\n").encode("utf-8"))
    fm = analyze(formula_b, formula_g, week="S39")
    fm_it = fm["rows"]["IT"]
    check(fm_it["mttr"]["P1"] == "04:00" and fm_it["mttr_ok"]["P1"] is True,
          f"MTTR P1 = (3+5)/2 = 4 h -> 04:00; 0/-4 excluded, open 99 h excluded: "
          f"{fm_it['mttr']['P1']}")
    check(fm_it["ouvert"] == {"P1": 1, "P2": 1},
          "Ouvert stays purely status-based (Duree ignored)")
    check(fm_it["mttr"]["P2"] == "02:00", f"P2 subset = F6 only -> 2/1 = {fm_it['mttr']['P2']}")
    check(fm["rows"]["Core CS"]["mttr"]["P2"] is None
          and fm["rows"]["Core CS"]["ouvert"]["P2"] == 1,
          "GNOC: blank status -> Ouvert even with a Duree, MTTR '-'")
    check(any("zero or negative" in w and "Duree <= 0" in w for w in fm["warnings"]),
          f"zero (F3) and negative (F4) Duree warned: {fm['warnings']}")

    # ---- single closed ticket in a domain -> its Duree IS the MTTR -------
    solo_b = ("b.csv", ("Ticket;semaine creation;identifiant EDS active;priorite traitement;"
                        "SLA;Duree\n"
                        "S1;S39;505010;P1;GOOD;5h30\n"
                        "S2;S39;505010;P2;En cours;\n").encode("utf-8"))
    solo_g = ("g.csv", ("Ticket;semaine creation;priorite traitement;valeur champ 2;"
                        "Statut SLA Retablissement;Duree\n"
                        "S3;S39;P1;core CS;GOOD;7,25\n").encode("utf-8"))
    sm = analyze(solo_b, solo_g, week="S39")
    check(sm["rows"]["IT"]["mttr"]["P1"] == "05:30",
          f"single closed ticket: MTTR = its Duree (5h30 -> 05:30): "
          f"{sm['rows']['IT']['mttr']['P1']}")
    check(sm["rows"]["IT"]["ouvert"] == {"P1": 0, "P2": 1},
          "single closed ticket -> MTTR, the open one -> Ouvert")
    check(sm["rows"]["Core CS"]["mttr"]["P1"] == "07:15",
          f"GNOC single closed ticket: MTTR = its Duree (7,25 -> 07:15): "
          f"{sm['rows']['Core CS']['mttr']['P1']}")

    # xlsx source with accented French headers
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Ticket", "Semaine création", "Priorité de traitement", "Valeur champ 2",
                  "SLA", "Date de création", "Date de résolution", "Durée",
                  "Identifiant EDS actif"])
    sheet.append(["X1", "S39", "P1", "core CS", "GOOD", "21/09/2026 08:00",
                  "21/09/2026 10:00", 2, "505010"])
    buf = io.BytesIO()
    book.save(buf)
    metrics_x = analyze(("tickets.xlsx", buf.getvalue()), GNOC, week="S39")
    check(metrics_x["columns"]["brise"]["week"] == "Semaine création", "xlsx accented week col")
    check(metrics_x["columns"]["brise"]["duration"] == "Durée", "xlsx accented duration col")
    check(metrics_x["rows"]["IT"]["tt"] == 1, "xlsx source row counted")
    check(metrics_x["rows"]["IT"]["mttr"]["P1"] == "02:00", "xlsx source MTTR")

    # xlsx empty status cell (None) = not evaluated -> open, no bogus warning;
    # and a single closed ticket with H:M:S '3:26:0' -> its Duree IS the MTTR
    book2 = Workbook()
    sh2 = book2.active
    sh2.append(["Ticket", "semaine", "Priorité traitement", "Valeur champ 2",
                "Statut SLA Retablissement", "Duree"])
    sh2.append(["N1", "S39", "P1", "core CS", None, "1:30"])
    sh2.append(["N2", "S39", "P1", "core CS", "GOOD", "3:26:0"])
    buf2 = io.BytesIO()
    book2.save(buf2)
    metrics_n = analyze(BRISE, ("gnoc_empty.xlsx", buf2.getvalue()), week="S39")
    check(metrics_n["rows"]["Core CS"]["ouvert"]["P1"] == 1,
          "xlsx empty status cell (None) -> open -> Ouvert")
    check(metrics_n["rows"]["Core CS"]["mttr"]["P1"] == "03:26",
          f"single closed ticket '3:26:0' -> MTTR 03:26: "
          f"{metrics_n['rows']['Core CS']['mttr']['P1']}")
    check(not any("unrecognised SLA" in w for w in metrics_n["warnings"]),
          f"empty status cell not flagged as unrecognised: {metrics_n['warnings']}")

    # ================= Excel rendering ====================================
    xlsx_name, xlsx = render_excel(metrics)
    check(xlsx_name == "hebdo_S39.xlsx", f"xlsx filename {xlsx_name}")
    wb = load_xlsx(xlsx)
    check(wb.sheetnames == ["Tickets", "MTTR P1", "MTTR P2"], f"sheets {wb.sheetnames}")
    check(wb.properties.title == "Rapport Hebdo S39", f"doc title {wb.properties.title}")

    tickets = wb["Tickets"]
    check(tickets.oddHeader.center.text == "Rapport Hebdo S39", "print header week")
    check("Orange Restricted" in (tickets.oddFooter.center.text or ""),
          "template footer preserved")

    # IT row = row 6: TT, rates, tallies
    check(tickets["B6"].value == 5, f"Tickets!B6 {tickets['B6'].value}")
    check(tickets["C6"].value == 0.5 and tickets["C6"].number_format == "0.00%",
          f"Tickets!C6 {tickets['C6'].value} fmt {tickets['C6'].number_format}")
    check(tickets["E6"].value == 1 and tickets["E6"].number_format == "0.00%", "Tickets!E6 100%")
    check(tickets["F6"].value == "-", f"Tickets!F6 {tickets['F6'].value!r}")
    check(tickets["G6"].value == 3 and tickets["H6"].value == 1 and tickets["I6"].value == 1,
          "Tickets!G6:I6 tallies")

    # Core PS (row 4) has no data at all -> dashes
    check(tickets["B4"].value == "-" and tickets["C4"].value == "-"
          and tickets["G4"].value == "-", "Core PS all dashes")
    # Core IP (row 3): evaluated 0% stays numeric, un-evaluated stays '-'
    check(tickets["E3"].value == 0 and tickets["E3"].number_format == "0.00%",
          f"Tickets!E3 real 0% {tickets['E3'].value}")
    check(tickets["F3"].value == "-", f"Tickets!F3 unevaluated P4 -> {tickets['F3'].value!r}")
    check(tickets["H3"].value == 1, f"Tickets!H3 Core IP Hors SLA {tickets['H3'].value!r}")
    check(tickets["I3"].value == 1, "Tickets!I3 Core IP en cours")
    # IN/TRANS/VAS (row 5): 2 GOOD of 3 P2 tickets
    check(abs(tickets["D5"].value - 2 / 3) < 1e-9, "Tickets!D5 IN/TRANS/VAS P2 rate")
    check(tickets["I5"].value == 0 and isinstance(tickets["I5"].value, int),
          "Tickets!I5 zero en cours on a row WITH tickets -> numeric 0")

    # MTTR P1 (cells are HH:MM text, colour-flagged against the 4 h target)
    mttr_p1 = wb["MTTR P1"]
    check(mttr_p1["B6"].value == "03:45", f"MTTR P1!B6 {mttr_p1['B6'].value!r}")
    check(mttr_p1["B6"].number_format == "General", f"MTTR P1!B6 fmt {mttr_p1['B6'].number_format}")
    check(cell_fill(mttr_p1["B6"]).upper().endswith("C6EFCE"),
          f"MTTR P1!B6 fill {cell_fill(mttr_p1['B6'])}")
    check(mttr_p1["C6"].value == 0 and isinstance(mttr_p1["C6"].value, int), "MTTR P1!C6 ouvert 0")
    check(mttr_p1["B2"].value == "01:00"
          and cell_fill(mttr_p1["B2"]).upper().endswith("C6EFCE"),
          "MTTR P1!B2 Core CS green")
    check(mttr_p1["C2"].value == 1, "MTTR P1!C2 Core CS ouvert")
    check(mttr_p1["B3"].value == "-" and cell_fill(mttr_p1["B3"]) == "",
          "MTTR P1!B3 resolved ticket without a valid Duree excluded -> '-' no fill")
    check(mttr_p1["C3"].value == 0 and isinstance(mttr_p1["C3"].value, int),
          "MTTR P1!C3 resolved status -> ouvert 0 (numeric)")
    check(mttr_p1["C4"].value == "-", "MTTR P1!C4 Core PS no P1 ticket -> '-'")

    # MTTR P2 - red flag when above the 8h target (sum/count = 30 h / 1 ticket)
    mttr_p2 = wb["MTTR P2"]
    check(mttr_p2["B5"].value == "30:00", f"MTTR P2!B5 {mttr_p2['B5'].value!r}")
    check(cell_fill(mttr_p2["B5"]).upper().endswith("FFC7CE"),
          f"MTTR P2!B5 red fill {cell_fill(mttr_p2['B5'])}")
    check(mttr_p2["B3"].value == "03:00"
          and cell_fill(mttr_p2["B3"]).upper().endswith("C6EFCE"), "MTTR P2!B3 green")
    check(mttr_p2["B6"].value == "01:30" and mttr_p2["C6"].value == 1, "MTTR P2!B6/C6 IT")
    check(mttr_p2["B2"].value == "-" and mttr_p2["C2"].value == "-",
          "MTTR P2!B2/C2 Core CS has no P2 ticket -> '-' / '-'")

    # dynamic naming for another week
    x40_name, x40 = render_excel(off)
    check(x40_name == "hebdo_S40.xlsx", f"off-week filename {x40_name}")
    wb40 = load_xlsx(x40)
    check(wb40["Tickets"]["B6"].value == "-", "off-week workbook is all dashes")
    check(wb40["MTTR P1"]["B6"].value == "-" and wb40["MTTR P1"]["C6"].value == "-",
          "off-week MTTR/Ouvert cells are '-' (no data)")
    check(wb40.properties.title == "Rapport Hebdo S40", "off-week doc title")

    # ================= PowerPoint rendering ===============================
    pptx_name, pptx = render_pptx(metrics)
    check(pptx_name == "hebdo_S39.pptx", f"pptx filename {pptx_name}")
    prs = load_pptx(pptx)
    check(len(prs.slides) == 4, f"pptx slides {len(prs.slides)}")
    check("Rapport Hebdo S39" in prs.slides[0].shapes.title.text, "pptx title slide")

    tickets_tbl = None
    for shape in prs.slides[1].shapes:
        if shape.has_table:
            tickets_tbl = shape.table
    check(tickets_tbl is not None and len(tickets_tbl.rows) == 6
          and len(tickets_tbl.columns) == 9, "pptx tickets table shape")
    last = [tickets_tbl.cell(5, c).text for c in range(9)]
    check(last == ["IT", "5", "50.00%", "50.00%", "100.00%", "-", "3", "1", "1"],
          f"pptx IT row {last}")

    red = green = 0
    red_font = green_font = 0
    for slide in (prs.slides[2], prs.slides[3]):
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            for r in range(1, len(table.rows)):
                cell = table.cell(r, 1)
                try:
                    rgb = str(cell.fill.fore_color.rgb or "")
                except Exception:  # noqa: BLE001
                    rgb = ""
                if rgb.upper().endswith("FFC7CE"):
                    red += 1
                elif rgb.upper().endswith("C6EFCE"):
                    green += 1
                # highlighted text: bold + red/green font on flagged cells
                para = cell.text_frame.paragraphs[0]
                try:
                    font_rgb = str(para.font.color.rgb or "")
                except Exception:  # noqa: BLE001
                    font_rgb = ""
                if para.font.bold:
                    if font_rgb.upper().endswith("BE123C"):
                        red_font += 1
                    elif font_rgb.upper().endswith("15803D"):
                        green_font += 1
    check(red == 1, f"pptx red MTTR cells {red}")
    check(green == 4, f"pptx green MTTR cells {green}")
    check(red_font == 1, f"pptx red bold MTTR fonts {red_font}")
    check(green_font == 4, f"pptx green bold MTTR fonts {green_font}")

    # ================= generate ===========================================
    files_both, _ = generate(BRISE, GNOC, week="S39")
    check([n for n, _ in files_both] == ["hebdo_S39.xlsx", "hebdo_S39.pptx"],
          f"generate both {[n for n, _ in files_both]}")
    files_x, _ = generate(BRISE, GNOC, week="S39", outputs=("xlsx",))
    check(len(files_x) == 1 and files_x[0][0] == "hebdo_S39.xlsx", "generate xlsx only")
    try:
        generate(BRISE, GNOC, week="S39", outputs=())
        check(False, "generate with no outputs should raise")
    except HebdoError:
        check(True, "")

    # ================= bundled sample data ================================
    from pathlib import Path

    sample_dir = Path(__file__).resolve().parent / "sample-data"
    brise_sample = sample_dir / "brise_smc_week39.csv"
    gnoc_sample = sample_dir / "GNOC_week39.csv"
    if brise_sample.exists() and gnoc_sample.exists():
        sample_metrics = analyze((brise_sample.name, brise_sample.read_bytes()),
                                 (gnoc_sample.name, gnoc_sample.read_bytes()), week="S39")
        check(sample_metrics["rows"]["IT"]["tt"] == 9, "sample brise: IT tt = 9")
        check(sample_metrics["rows"]["IT"]["mttr"]["P1"] == "04:10",
              "sample brise: MTTR P1 (avg 4.17 h -> 04:10)")
        check(sample_metrics["rows"]["IT"]["mttr_ok"]["P1"] is False,
              "sample brise: 4.17 h above the 4 h target -> BAD")
        check(sample_metrics["rows"]["IT"]["mttr"]["P2"] == "33:00",
              "sample brise: MTTR P2 shown beyond 24 h (33:00)")
        check(sample_metrics["rows"]["IT"]["mttr_ok"]["P2"] is False,
              "sample brise: 33 h above the 8 h target -> BAD")
        check(sample_metrics["rows"]["Core PS"]["mttr_ok"]["P1"] is False,
              "sample GNOC: Core PS P1 above 4h target")
        check(sample_metrics["rows"]["IN/TRANS/VAS"]["tt"] == 4,
              "sample GNOC: IN + VAS + TRANS + literal all in IN/TRANS/VAS")
        check(sample_metrics["rows"]["Core IP"]["rates"]["P3"] == 0.0,
              "sample: evaluated zero stays numeric 0.00%")
        check(sample_metrics["rows"]["Core IP"]["rates"]["P4"] is None,
              "sample: unevaluated SLA stays empty ('-')")
        # status rules on the shipped samples (GNOC header renamed)
        check(sample_metrics["rows"]["IT"]["ouvert"] == {"P1": 0, "P2": 1},
              "sample brise: 'En Cours' ticket counted Ouvert")
        check(sample_metrics["rows"]["Core CS"]["ouvert"]["P1"] == 1
              and sample_metrics["rows"]["Core CS"]["ouvert"]["P2"] == 0,
              "sample GNOC: blank status = Ouvert P1, resolved rows = 0")
    else:
        check(False, "bundled sample files missing")


if __name__ == "__main__":
    run_tests()
    if FAILURES:
        print(f"FAILED {len(FAILURES)}/{CHECKS} checks:")
        for failure in FAILURES:
            if failure:
                print(f"  - {failure}")
        raise SystemExit(1)
    print(f"OK: all {CHECKS} checks passed")
