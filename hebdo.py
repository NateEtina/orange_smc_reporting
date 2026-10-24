"""Weekly Excel & PowerPoint report generation - Feature 1 (P1/P2 MTTR).

Ingests the two weekly sources (brise_smc_weekXX + GNOC_weekXX), applies the
business rules of the weekly "Hebdo" report and fills the template workbook
(templates/hebdo_S38.xlsx -> output hebdo_SXX.xlsx), plus a generated
PowerPoint deck (hebdo_SXX.pptx).

Business rules implemented here
-------------------------------
* Week management   : explicit week label ("S39"/"39") or date -> ISO week;
                      all filtering uses the source column "semaine creation".
* brise_smc         : keep only active EDS units 505010 / 555402 / 610136,
                      inject into the "IT" row of the 3 template sheets.
* GNOC              : group by "valeur champ 2" into the rows Core CS,
                      Core IP, Core PS, IN/TRANS/VAS. IN/TRANS/VAS is a
                      composite domain: the cells IN, VAS and TRANS (and any
                      combined form such as IN/TRANS/VAS) all map to it.
* Tickets sheet     : TT count, per-priority SLA rate (GOOD / total of that
                      priority). A priority whose tickets carry no evaluated
                      SLA verdict at all has *no data* -> "-", which is
                      distinct from a real calculated 0.00%.
                      Tallies: GOOD / BAD / blank (In SLA / Hors SLA / En cours).
* MTTR sheets       : per priority (P1, P2) two columns driven by the SLA
                      status verdict ("SLA" in brise, "Statut SLA Retablissement"
                      in GNOC), case-insensitive & trimmed:
                        - brise: OPEN when the status equals "En cours";
                                 resolved otherwise.
                        - GNOC : OPEN when the status is blank;
                                 resolved otherwise.
                      "Ouvert" counts the OPEN tickets of that priority.
                      MTTR = SUM(Duree) / COUNT over the subset of tickets
                      that are, simultaneously: of that priority, of that
                      domain, RESOLVED (per the file rule above) and with a
                      valid duration Duree > 0 (strict). Resolved tickets
                      whose Duree is blank, unparsable or <= 0 are excluded
                      from both the sum and the count (warned); when the
                      subset is empty the cell shows "-". When exactly ONE
                      closed ticket qualifies for a domain, its Duree IS the
                      MTTR (sum 1 / count 1).
                      The value is expressed in total hours and displayed as
                      HH:MM (36.5 h -> "36:30", works beyond 24 h);
                      "-" when no ticket qualifies.
                      Cells are colour-flagged against the SLA target:
                      P1 > 4.0 h and P2 > 8.0 h are BAD (red), otherwise OK.
* Display convention: a calculated zero is shown as 0 (0.00% for rates);
                      cells with no underlying data (empty domain, no ticket
                      of that priority) show "-".
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from converter import (
    decode_text,
    detect_delimiter,
    read_rows,
    sanitize_header,
)

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "templates" / "hebdo_S38.xlsx"

# --------------------------------------------------------------------------- #
# Business constants
# --------------------------------------------------------------------------- #

EDS_ACTIVE = ("505010", "555402", "610136")
TARGET_HOURS: Dict[str, float] = {"P1": 4.0, "P2": 8.0}
PRIORITIES = ("P1", "P2", "P3", "P4")
MTTR_PRIORITIES = ("P1", "P2")

TICKET_ROWS = ("Core CS", "Core IP", "Core PS", "IN/TRANS/VAS", "IT")
GNOC_ROWS = TICKET_ROWS[:-1]
IT_ROW = "IT"

SHEET_TICKETS = "Tickets"
SHEET_MTTR = {"P1": "MTTR P1", "P2": "MTTR P2"}
TICKET_COLUMNS = {"tt": 2, "P1": 3, "P2": 4, "P3": 5, "P4": 6,
                  "in_sla": 7, "hors_sla": 8, "en_cours": 9}

PCT_FORMAT = "0.00%"
FILL_OK = "C6EFCE"      # green - MTTR within target
FILL_BAD = "FFC7CE"     # red   - MTTR above target
FONT_OK = "15803d"      # dark green text - PowerPoint MTTR highlight
FONT_BAD = "be123c"     # dark red text   - PowerPoint MTTR highlight

# --------------------------------------------------------------------------- #
# Source column roles
# --------------------------------------------------------------------------- #

ROLES = ("week", "eds", "priority", "sla", "domain", "duration")
DEFAULT_COLUMNS: Dict[str, str] = {
    "week": "semaine creation",
    "eds": "identifiant EDS active",
    "priority": "priorite traitement",
    "sla": "SLA",
    "domain": "valeur champ 2",
    "duration": "Duree",
}
REQUIRED_BRISE = ("week", "eds", "priority", "sla", "duration")
REQUIRED_GNOC = ("week", "priority", "sla", "domain", "duration")
OVERRIDABLE = {
    "brise": ("week", "eds", "priority", "sla", "duration"),
    "gnoc": ("week", "priority", "sla", "domain", "duration"),
}

_DOMAIN_KEYS = {
    "corecs": "Core CS",
    "coreip": "Core IP",
    "coreps": "Core PS",
    "intransvas": "IN/TRANS/VAS",
}
# The composite IN/TRANS/VAS row also swallows its 3 distinct sub-domains.
COMPOSITE_PARTS = frozenset({"in", "vas", "trans"})

WARN_SAMPLE = 5          # max distinct offending values listed in a warning


class HebdoError(Exception):
    """A problem the user can fix: bad week, missing column, broken template."""


# =========================================================================== #
# Week management
# =========================================================================== #

_WEEK_PLAIN = re.compile(r"^(\d{1,2})(?:\.0+)?$")
_WEEK_LABEL = re.compile(r"^(?:SEMAINE|SEM|S|W)\s*[-.]?\s*(\d{1,2})$", re.I)
_WEEK_YEARED = re.compile(r"^(\d{4})\s*[-/]\s*(?:SEMAINE|SEM|S|W)?\s*(\d{1,2})$", re.I)
_DATE_TEXT_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y",
                      "%Y%m%d", "%d/%m/%y", "%m/%d/%Y")


def parse_week(value: Any) -> int:
    """Accept 39, "39", "S39", "s 39", "Semaine 39" -> 39."""
    text = str(value).strip()
    if not text:
        raise HebdoError("Week is empty.")
    m = _WEEK_LABEL.match(text) or _WEEK_PLAIN.match(text)
    if not m:
        raise HebdoError(f"Cannot read a week number from {value!r} (expected e.g. 'S39' or '39').")
    week = int(m.group(1))
    if not 1 <= week <= 53:
        raise HebdoError(f"Invalid week number: {week} (must be 1-53).")
    return week


def parse_day(value: Any) -> date:
    """Accept 2026-09-21 / 21-09-2026 / 21/09/2026 -> date."""
    text = str(value).strip()
    if not text:
        raise HebdoError("Date is empty.")
    for fmt in _DATE_TEXT_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        raise HebdoError(f"Cannot read a date from {value!r} (expected e.g. 2026-09-21).") from None


def resolve_week(week: Any = None, date_str: Any = None) -> Tuple[int, int]:
    """Return (iso_week, iso_year). An explicit week wins over the date; a date
    wins over 'today'."""
    if week not in (None, ""):
        number = parse_week(week)
        if date_str not in (None, ""):
            year = parse_day(date_str).isocalendar().year
        else:
            year = date.today().isocalendar().year
        return number, year
    day = parse_day(date_str) if date_str not in (None, "") else date.today()
    iso = day.isocalendar()
    return iso[1], iso[0]


def week_label(week: int) -> str:
    return f"S{week}"


def week_bounds(week: int, year: int) -> Tuple[date, date]:
    return (date.fromisocalendar(year, week, 1), date.fromisocalendar(year, week, 7))


def week_value_matches(value: Any, week: int, year: int) -> bool:
    """True when a 'semaine creation' cell designates the target week."""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        iso = value.isocalendar()
        return iso.week == week and iso.year == year
    if isinstance(value, (int, float)):
        return float(value).is_integer() and int(value) == week

    text = str(value).strip()
    if not text:
        return False
    m = _WEEK_PLAIN.match(text)
    if m:
        return int(m.group(1)) == week
    m = _WEEK_LABEL.match(text)
    if m:
        return int(m.group(1)) == week
    m = _WEEK_YEARED.match(text)
    if m:
        return int(m.group(2)) == week and int(m.group(1)) == year

    for candidate in (text, text[:10]):
        for fmt in _DATE_TEXT_FORMATS:
            try:
                day = datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
            iso = day.isocalendar()
            return iso.week == week and iso.year == year
    return False


# =========================================================================== #
# Reading sources (CSV / TSV / TXT / XLSX)
# =========================================================================== #

def read_records(filename: str, raw: bytes) -> Tuple[List[str], List[Dict[str, Any]], List[str]]:
    """Return (headers, records, warnings) for either CSV text or XLSX bytes."""
    if not raw:
        raise HebdoError(f"'{filename}' is empty.")
    warnings: List[str] = []
    lower = (filename or "").lower()
    if lower.endswith((".xlsx", ".xlsm")) or raw.startswith(b"PK\x03\x04"):
        return _read_xlsx(filename, raw, warnings)
    return _read_csv(filename, raw, warnings)


def _read_csv(filename: str, raw: bytes, warnings: List[str]):
    text, _encoding, _conf, decode_warnings = decode_text(raw)
    warnings.extend(decode_warnings)
    delimiter = detect_delimiter(text)
    rows = [r for r in read_rows(text, delimiter, '"')
            if any((cell or "").strip() for cell in r)]
    if not rows:
        raise HebdoError(f"'{filename}' has no header row.")
    header = sanitize_header(rows[0])
    records: List[Dict[str, Any]] = []
    for row in rows[1:]:
        records.append({header[i]: (row[i] if i < len(row) else "") for i in range(len(header))})
    if not records:
        raise HebdoError(f"'{filename}' contains a header but no data rows.")
    return header, records, warnings


def _read_xlsx(filename: str, raw: bytes, warnings: List[str]):
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface any reader failure as a business error
        raise HebdoError(f"Could not read '{filename}' as an Excel file: {exc}") from exc
    try:
        sheet = book.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        first = next(rows_iter, None)
        while first is not None and not any(v is not None and str(v).strip() for v in first):
            first = next(rows_iter, None)
        if first is None:
            raise HebdoError(f"'{filename}' has no header row.")
        header = sanitize_header(["" if v is None else str(v) for v in first])
        records: List[Dict[str, Any]] = []
        for row in rows_iter:
            if row is None or not any(v is not None and str(v).strip() for v in row):
                continue
            records.append({header[i]: (row[i] if i < len(row) else "")
                            for i in range(len(header))})
        if not records:
            raise HebdoError(f"'{filename}' contains a header but no data rows.")
        return header, records, warnings
    finally:
        book.close()


# =========================================================================== #
# Column resolution
# =========================================================================== #

def _fold(value: Any) -> str:
    """Accent-stripped, lower-cased, whitespace-collapsed key for a header."""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", text).strip().casefold().rstrip(":").strip()


def _contains(*needles: str):
    return lambda header: all(n in header for n in needles)


_ROLE_TESTS: Dict[str, Tuple] = {
    "week": (
        lambda h: h == "semaine creation",
        _contains("semaine", "creation"),
        lambda h: h == "semaine",
        lambda h: h == "week",
    ),
    "eds": (
        lambda h: h == "identifiant eds active",
        _contains("eds", "active"),
        lambda h: h == "eds",
        _contains("eds",),
    ),
    "priority": (
        lambda h: h == "priorite traitement",
        lambda h: h in ("priorite", "priority", "prio"),
        lambda h: "priorite" in h or "priority" in h,
    ),
    "sla": (
        lambda h: h == "sla",
        lambda h: h == "statut sla retablissement",
        lambda h: "sla" in h and "hors" not in h,
    ),
    "domain": (
        lambda h: h == "valeur champ 2",
        lambda h: "valeur champ" in h and "2" in h,
        lambda h: "champ 2" in h or "champ2" in h,
        lambda h: h in ("domaine", "domain", "domaine 2"),
        lambda h: "domaine" in h or "domain" in h,
    ),
    "duration": (
        lambda h: h == "duree",
        lambda h: h.startswith("duree"),
        lambda h: "duree" in h,
        lambda h: ("resolution time" in h or "temps de resolution" in h) and "date" not in h,
    ),
}


def resolve_columns(headers: Sequence[str], overrides: Optional[Dict[str, str]] = None,
                    required: Sequence[str] = (),
                    roles: Sequence[str] = ROLES) -> Dict[str, Optional[str]]:
    """Map each logical role to a concrete header (explicit override wins,
    then keyword detection)."""
    overrides = overrides or {}
    folded = {h: _fold(h) for h in headers}
    resolved: Dict[str, Optional[str]] = {}
    used: set = set()

    for role in roles:
        want = overrides.get(role)
        if not want:
            continue
        hit = next((h for h in headers if h == want or folded[h] == _fold(want)), None)
        if hit is None:
            raise HebdoError(
                f"Column '{want}' (for '{role}') not found. Available columns: "
                + ", ".join(headers)
            )
        resolved[role] = hit
        used.add(hit)

    for role in roles:
        if role in resolved:
            continue
        found: Optional[str] = None
        for test in _ROLE_TESTS[role]:
            found = next((h for h in headers if h not in used and test(folded[h])), None)
            if found:
                break
        resolved[role] = found
        if found:
            used.add(found)

    missing = [role for role in required if not resolved.get(role)]
    if missing:
        # The SLA verdict is spelled 'SLA' in brise_smc and
        # 'Statut SLA Retablissement' in GNOC exports.
        expected = {"sla": "SLA / Statut SLA Retablissement"}
        raise HebdoError(
            "Missing required column(s): "
            + ", ".join(f"{role} (expected '{expected.get(role, DEFAULT_COLUMNS[role])}')"
                        for role in missing)
            + ". Found: " + ", ".join(headers)
        )
    return resolved


# =========================================================================== #
# Value normalisation
# =========================================================================== #

def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _eds_id(value: Any) -> str:
    text = re.sub(r"\s+", "", _cell(value))
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return text


def _priority(value: Any) -> Optional[str]:
    text = _fold(value).replace(" ", "")
    m = re.fullmatch(r"p?([1-4])(?:eme|ème|e)?", text)
    return f"P{m.group(1)}" if m else None


def _sla_status(value: Any) -> Tuple[str, bool]:
    """Return (status, recognised). status is 'good' | 'bad' | 'open'."""
    if value is None or not _cell(value):
        # Empty cell (GNOC exports write blanks as None) = not evaluated yet.
        return "open", True
    text = _fold(value)
    if not text or text == "vide":
        return "open", True
    if text in ("good", "ok", "in sla", "en sla"):
        return "good", True
    if text in ("bad", "ko", "hors sla"):
        return "bad", True
    if text.startswith("en cours") or text in ("encours", "en-cours", "in progress",
                                               "inprogress", "wip"):
        return "open", True
    if "good" in text:
        return "good", True
    if "bad" in text:
        return "bad", True
    return "open", False


def _domain_row(value: Any) -> Optional[str]:
    """Map a 'valeur champ 2' cell to a report row.

    IN, VAS and TRANS are the 3 distinct members of the composite
    IN/TRANS/VAS row, so each of them (and any combined spelling such as
    'IN/TRANS/VAS') maps there. Matching is token-based, so an unrelated
    domain like 'INCIDENT' cannot be mistaken for 'IN'.
    """
    text = _fold(value)
    if not text:
        return None
    compact = re.sub(r"[^a-z0-9]+", "", text)
    if compact == "intransvas":
        return "IN/TRANS/VAS"
    tokens = [t for t in re.split(r"[^a-z0-9]+", text) if t]
    if tokens and any(t in COMPOSITE_PARTS for t in tokens):
        return "IN/TRANS/VAS"
    return _DOMAIN_KEYS.get(compact)


# =========================================================================== #
# Duration parsing (MTTR: the 'Duree' resolution-time column)
# =========================================================================== #

_DURATION_TIME = re.compile(r"^(\d{1,4}):([0-5]?\d)(?::([0-5]?\d))?$")
# ^ lenient on purpose: minutes/seconds may be 1 digit ('3:26:0', '0:6:59'),
#   which is how the brise/GNOC exports spell their H:M:S durations.
_DURATION_NUMBER = re.compile(r"^[+-]?\d+(?:[.,]\d+)?$")
_DURATION_SUFFIX = re.compile(r"^([+-]?\d+(?:[.,]\d+)?)\s*h(?:eures?)?$", re.I)
# compact French hours+minutes: '5h30', '5H30', '5 h 30', '12h45m' = 12.75 h
_DURATION_HM = re.compile(r"^(\d{1,4})\s*h\s*([0-5]?\d)(?:\s*m(?:in(?:ute)?s?)?)?$", re.I)


def parse_duration(value: Any) -> Optional[float]:
    """Resolution time in hours. ``None`` = blank or unrecognised.

    Accepted shapes: plain numbers (hours: 36.5 or 36,5), ``HH:MM`` /
    ``HH:MM:SS`` strings (36:30 = 36.5 h, works beyond 24 h) with lenient
    single-digit minutes/seconds as exported by brise/GNOC ('3:26:0',
    '0:6:59'), compact French hours+minutes (``5h30`` / ``5 h 30`` /
    ``12h45m``), a trailing ``h`` ('4.5 h'), Excel time cells
    (``datetime.time``) and Excel duration cells (``timedelta``).
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, timedelta):
        return value.total_seconds() / 3600.0
    if isinstance(value, time):
        return value.hour + value.minute / 60.0 + value.second / 3600.0
    if isinstance(value, (int, float)):
        return float(value)
    text = _cell(value)
    if not text:
        return None
    m = _DURATION_TIME.match(text)
    if m:
        hours = int(m.group(1))
        minutes = int(m.group(2))
        seconds = int(m.group(3)) if m.group(3) else 0
        return hours + minutes / 60.0 + seconds / 3600.0
    m = _DURATION_HM.match(text)
    if m:
        return int(m.group(1)) + int(m.group(2)) / 60.0
    m = _DURATION_SUFFIX.match(text)
    if m:
        return float(m.group(1).replace(",", "."))
    if _DURATION_NUMBER.match(text):
        return float(text.replace(",", "."))
    return None


def format_mttr(avg_hours: float) -> str:
    """36.5 h -> '36:30' (total hours:minutes; values beyond 24 h kept)."""
    total_minutes = int(avg_hours * 60 + 0.5)
    hours, minutes = divmod(total_minutes, 60)
    return f"{hours:02d}:{minutes:02d}"


# =========================================================================== #
# Aggregation
# =========================================================================== #

@dataclass
class RowStats:
    tt: int = 0
    in_sla: int = 0
    hors_sla: int = 0
    en_cours: int = 0
    prio_total: Dict[str, int] = field(default_factory=lambda: {p: 0 for p in PRIORITIES})
    prio_good: Dict[str, int] = field(default_factory=lambda: {p: 0 for p in PRIORITIES})
    prio_bad: Dict[str, int] = field(default_factory=lambda: {p: 0 for p in PRIORITIES})
    durations: Dict[str, List[float]] = field(default_factory=lambda: {p: [] for p in MTTR_PRIORITIES})
    open_counts: Dict[str, int] = field(default_factory=lambda: {p: 0 for p in MTTR_PRIORITIES})


def _sample(values: set) -> List[str]:
    return sorted(str(v) for v in list(values)[:WARN_SAMPLE])


def _aggregate(records: Sequence[Dict[str, Any]], columns: Dict[str, Optional[str]],
               *, week: int, year: int, fixed_row: Optional[str], eds_filter: bool,
               source: str
               ) -> Tuple[Dict[str, RowStats], Dict[str, Any]]:
    """source: 'brise' or 'gnoc' - selects the open/resolved status rule."""
    stats: Dict[str, RowStats] = {}
    week_col = columns["week"]
    prio_col = columns["priority"]
    sla_col = columns["sla"]
    eds_col = columns.get("eds")
    domain_col = columns.get("domain")
    duration_col = columns.get("duration")

    info: Dict[str, Any] = {
        "rows": len(records),
        "matched_week": 0,
        "excluded_eds": 0,
        "excluded_domain": 0,
        "excluded_priority": 0,
        "negative_durations": 0,
        "unparsed_durations": 0,
        "blank_durations": 0,
        "unknown_priority": set(),
        "unknown_domain": set(),
        "unknown_sla": set(),
    }

    for rec in records:
        if not week_value_matches(rec.get(week_col), week, year):
            continue
        info["matched_week"] += 1
        if eds_filter and _eds_id(rec.get(eds_col)) not in EDS_ACTIVE:
            info["excluded_eds"] += 1
            continue
        priority = _priority(rec.get(prio_col))
        if priority is None:
            info["excluded_priority"] += 1
            info["unknown_priority"].add(_cell(rec.get(prio_col)) or "(empty)")
            continue
        row = fixed_row or (_domain_row(rec.get(domain_col)) if domain_col else None)
        if row is None:
            info["excluded_domain"] += 1
            info["unknown_domain"].add(_cell(rec.get(domain_col)) or "(empty)")
            continue

        st = stats.setdefault(row, RowStats())
        st.tt += 1
        st.prio_total[priority] += 1

        status, known = _sla_status(rec.get(sla_col))
        if not known:
            info["unknown_sla"].add(_cell(rec.get(sla_col)))
        if status == "good":
            st.in_sla += 1
            st.prio_good[priority] += 1
        elif status == "bad":
            st.hors_sla += 1
            st.prio_bad[priority] += 1
        else:
            st.en_cours += 1

        # Open vs resolved (file-specific status rule on the SLA verdict:
        # 'SLA' in brise, 'Statut SLA Retablissement' in GNOC):
        #   brise : open <=> status equals "En cours"   (case-insensitive)
        #   GNOC  : open <=> status is blank
        # MTTR subset = RESOLVED tickets with Duree > 0 (strict): the formula
        # is sum(Duree) / count of that subset. A resolved ticket whose Duree
        # is blank, unparsable or <= 0 is excluded from both (warned below).
        if priority in MTTR_PRIORITIES:
            status_text = _cell(rec.get(sla_col))
            is_open = ((status_text == "") if source == "gnoc"
                       else (_fold(status_text) == "en cours"))
            if is_open:
                st.open_counts[priority] += 1
            elif duration_col:
                raw_duration = rec.get(duration_col)
                blank = raw_duration is None or not _cell(raw_duration)
                hours = None if blank else parse_duration(raw_duration)
                if hours is None and not blank:
                    info["unparsed_durations"] += 1
                if hours is not None and hours <= 0:
                    # 0 h counts with negatives: the MTTR rule is Duree > 0 (warned)
                    info["negative_durations"] += 1
                if blank:
                    info["blank_durations"] += 1
                if hours is not None and hours > 0:
                    st.durations[priority].append(hours)

    return stats, info


def _mttr_of(st: RowStats, priority: str) -> Tuple[Optional[str], Optional[bool]]:
    """MTTR = sum(Duree) / count over the filtered closed tickets of that
    priority (Duree > 0 only). Single-ticket rule: when exactly ONE closed
    ticket qualifies for the domain, its Duree IS the MTTR (sum 1 / count 1).
    No qualifying ticket -> no data ('-')."""
    values = st.durations[priority]
    if not values:
        return None, None
    avg = values[0] if len(values) == 1 else sum(values) / len(values)
    return format_mttr(avg), avg <= TARGET_HOURS[priority]


def _row_metrics(st: Optional[RowStats]) -> Dict[str, Any]:
    st = st or RowStats()
    rates: Dict[str, Optional[float]] = {}
    for prio in PRIORITIES:
        total = st.prio_total[prio]
        evaluated = st.prio_good[prio] + st.prio_bad[prio]
        # No ticket at all, or none with an evaluated verdict -> no data ("->"),
        # which stays distinct from a calculated 0.00% rate.
        rates[prio] = (st.prio_good[prio] / total) if total and evaluated else None
    mttr: Dict[str, Optional[str]] = {}
    mttr_ok: Dict[str, Optional[bool]] = {}
    ouvert: Dict[str, Optional[int]] = {}
    for prio in MTTR_PRIORITIES:
        mttr[prio], mttr_ok[prio] = _mttr_of(st, prio)
        # No ticket of that priority -> no data ("-"), not a calculated 0.
        ouvert[prio] = st.open_counts[prio] if st.prio_total[prio] else None
    # Zero vs no-data: an empty domain row yields "-" in every count cell,
    # while a row with tickets keeps its numeric counts (0 included).
    empty = st.tt == 0
    return {
        "tt": None if empty else st.tt,
        "rates": rates,
        "in_sla": None if empty else st.in_sla,
        "hors_sla": None if empty else st.hors_sla,
        "en_cours": None if empty else st.en_cours,
        "prio_total": dict(st.prio_total),
        "mttr": mttr,
        "mttr_ok": mttr_ok,
        "ouvert": ouvert,
    }


def _source_info(filename: str, info: Dict[str, Any], warnings: List[str]) -> Dict[str, Any]:
    unknown_prio = info["unknown_priority"]
    if unknown_prio:
        warnings.append(f"{filename}: ignored {info['excluded_priority']} row(s) with an "
                        f"unknown priority: {', '.join(_sample(unknown_prio))}")
    if info["excluded_domain"]:
        warnings.append(f"{filename}: ignored {info['excluded_domain']} row(s) outside the "
                        f"known sub-domains: {', '.join(_sample(info['unknown_domain']))}")
    if info["excluded_eds"]:
        warnings.append(f"{filename}: ignored {info['excluded_eds']} row(s) whose active EDS "
                        f"unit is not in {', '.join(EDS_ACTIVE)}.")
    if info["unknown_sla"]:
        warnings.append(f"{filename}: unrecognised SLA value(s) counted as 'En cours': "
                        f"{', '.join(_sample(info['unknown_sla']))}")
    if info["negative_durations"]:
        warnings.append(f"{filename}: {info['negative_durations']} zero or negative resolution "
                        f"time(s) (Duree <= 0) - excluded from MTTR (Duree > 0 required).")
    if info["unparsed_durations"]:
        warnings.append(f"{filename}: {info['unparsed_durations']} resolution time value(s) could "
                        f"not be parsed - excluded from MTTR (Duree > 0 required).")
    if info["blank_durations"]:
        warnings.append(f"{filename}: {info['blank_durations']} resolved ticket(s) without a "
                        f"resolution time - excluded from MTTR (Duree > 0 required).")
    if info["matched_week"] == 0:
        warnings.append(f"{filename}: no row matches the target week - check the week column.")
    return {
        "file": filename,
        "rows": info["rows"],
        "matched_week": info["matched_week"],
        "used": info["matched_week"] - info["excluded_eds"] - info["excluded_domain"]
                - info["excluded_priority"],
        "excluded_eds": info["excluded_eds"],
        "excluded_domain": info["excluded_domain"],
        "excluded_priority": info["excluded_priority"],
    }


# =========================================================================== #
# Public API
# =========================================================================== #

def analyze(brise: Tuple[str, bytes], gnoc: Tuple[str, bytes],
            *, week: Any = None, date: Any = None,
            columns: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    """Process both sources and return the JSON-serialisable report metrics.

    ``brise``/``gnoc`` are ``(filename, raw_bytes)`` tuples.
    """
    week_no, year = resolve_week(week, date)
    columns = columns or {}
    warnings: List[str] = []

    b_headers, b_records, b_warn = read_records(*brise)
    g_headers, g_records, g_warn = read_records(*gnoc)
    warnings.extend(f"brise: {w}" for w in b_warn)
    warnings.extend(f"GNOC: {w}" for w in g_warn)

    b_cols = resolve_columns(b_headers, columns, REQUIRED_BRISE, OVERRIDABLE["brise"])
    g_cols = resolve_columns(g_headers, columns, REQUIRED_GNOC, OVERRIDABLE["gnoc"])

    b_stats, b_info = _aggregate(b_records, b_cols, week=week_no, year=year,
                                 fixed_row=IT_ROW, eds_filter=True, source="brise")
    g_stats, g_info = _aggregate(g_records, g_cols, week=week_no, year=year,
                                 fixed_row=None, eds_filter=False, source="gnoc")

    sources = [
        _source_info(_display_name(brise[0]), b_info, warnings),
        _source_info(_display_name(gnoc[0]), g_info, warnings),
    ]

    rows: Dict[str, Any] = {}
    for label in TICKET_ROWS:
        stats = b_stats.get(label) if label == IT_ROW else g_stats.get(label)
        rows[label] = _row_metrics(stats)

    start, end = week_bounds(week_no, year)
    return {
        "week": week_no,
        "week_label": week_label(week_no),
        "year": year,
        "week_start": start.isoformat(),
        "week_end": end.isoformat(),
        "row_order": list(TICKET_ROWS),
        "rows": rows,
        "columns": {
            "brise": {role: b_cols.get(role) for role in ROLES},
            "gnoc": {role: g_cols.get(role) for role in ROLES},
        },
        "sources": sources,
        "warnings": warnings,
    }


def _display_name(filename: str) -> str:
    return (filename or "source").replace("\\", "/").split("/")[-1]


# =========================================================================== #
# Excel output
# =========================================================================== #

def _label_row_map(sheet, labels: Sequence[str]) -> Dict[str, int]:
    wanted = {_fold(label): label for label in labels}
    found: Dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        key = _fold(sheet.cell(row=row, column=1).value or "")
        if key in wanted:
            found[key] = row
    missing = [label for label in labels if _fold(label) not in found]
    if missing:
        raise HebdoError(
            f"Template sheet '{sheet.title}' is missing expected row(s): {', '.join(missing)}."
        )
    return {_fold(label): found[_fold(label)] for label in labels}


def _set_count(cell, count: Optional[int]) -> None:
    """Counts: a real number when the row has data (0 included), '-' otherwise."""
    if count is None:
        cell.value = "-"
    else:
        cell.value = int(count)
    cell.number_format = "General"
    cell.fill = PatternFill()


def _set_rate(cell, rate: Optional[float]) -> None:
    if rate is None:
        cell.value = "-"
        cell.number_format = "General"
        cell.fill = PatternFill()
    else:
        cell.value = rate
        cell.number_format = PCT_FORMAT
        cell.fill = PatternFill()


def _set_mttr(cell, text: Optional[str], ok: Optional[bool]) -> None:
    cell.value = text or "-"
    cell.number_format = "General"
    if text and ok is True:
        cell.fill = PatternFill("solid", fgColor=FILL_OK)
    elif text and ok is False:
        cell.fill = PatternFill("solid", fgColor=FILL_BAD)
    else:
        cell.fill = PatternFill()


def render_excel(metrics: Dict[str, Any],
                 template: Optional[Path] = None) -> Tuple[str, bytes]:
    """Fill the hebdo template with the metrics -> (filename, xlsx bytes)."""
    path = Path(template) if template else TEMPLATE_PATH
    if not path.exists():
        raise HebdoError(f"Report template not found: {path}")
    try:
        book = load_workbook(path)
    except Exception as exc:  # noqa: BLE001
        raise HebdoError(f"Could not open the report template: {exc}") from exc

    label = metrics["week_label"]
    rows = metrics["rows"]

    sheet = book[SHEET_TICKETS] if SHEET_TICKETS in book.sheetnames else None
    if sheet is None:
        raise HebdoError(f"Template has no '{SHEET_TICKETS}' sheet.")
    ticket_rows = _label_row_map(sheet, TICKET_ROWS)
    for name in TICKET_ROWS:
        row = ticket_rows[_fold(name)]
        m = rows[name]
        _set_count(sheet.cell(row=row, column=TICKET_COLUMNS["tt"]), m["tt"])
        for offset, prio in enumerate(PRIORITIES):
            _set_rate(sheet.cell(row=row, column=TICKET_COLUMNS[prio]), m["rates"][prio])
        _set_count(sheet.cell(row=row, column=TICKET_COLUMNS["in_sla"]), m["in_sla"])
        _set_count(sheet.cell(row=row, column=TICKET_COLUMNS["hors_sla"]), m["hors_sla"])
        _set_count(sheet.cell(row=row, column=TICKET_COLUMNS["en_cours"]), m["en_cours"])

    for prio in MTTR_PRIORITIES:
        sheet_name = SHEET_MTTR[prio]
        mttr_sheet = book[sheet_name] if sheet_name in book.sheetnames else None
        if mttr_sheet is None:
            raise HebdoError(f"Template has no '{sheet_name}' sheet.")
        mttr_rows = _label_row_map(mttr_sheet, TICKET_ROWS)
        for name in TICKET_ROWS:
            row = mttr_rows[_fold(name)]
            m = rows[name]
            _set_mttr(mttr_sheet.cell(row=row, column=2), m["mttr"][prio], m["mttr_ok"][prio])
            ouvert = m["ouvert"][prio]
            cell = mttr_sheet.cell(row=row, column=3)
            cell.value = ouvert if ouvert is not None else "-"
            cell.number_format = "General"
            cell.fill = PatternFill()

    # Dynamic week markers: file name, document properties, print header.
    book.properties.title = f"Rapport Hebdo {label}"
    book.properties.subject = (f"Tickets & Incidents {label} "
                               f"({metrics['week_start']} -> {metrics['week_end']})")
    for ws in book.worksheets:
        ws.oddHeader.center.text = f"Rapport Hebdo {label}"

    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return f"hebdo_{label}.xlsx", buffer.getvalue()


# =========================================================================== #
# PowerPoint output
# =========================================================================== #

def render_pptx(metrics: Dict[str, Any]) -> Tuple[str, bytes]:
    """Generate the PowerPoint deck from scratch -> (filename, pptx bytes)."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    label = metrics["week_label"]
    rows = metrics["rows"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_table(slide, headers, data, top, col_widths, fills=None, fonts=None, font_size=12):
        shape = slide.shapes.add_table(len(data) + 1, len(headers),
                                       Inches(0.5), Inches(top),
                                       Inches(12.333), Inches(0.4 + 0.36 * (len(data) + 1)))
        table = shape.table
        for i, width in enumerate(col_widths):
            table.columns[i].width = Inches(width)
        for c, text in enumerate(headers):
            cell = table.cell(0, c)
            cell.text = str(text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x78)
            para = cell.text_frame.paragraphs[0]
            para.font.bold = True
            para.font.size = Pt(font_size)
            para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for r, row in enumerate(data, start=1):
            for c, text in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(text)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(font_size)
                if fills and (r - 1, c) in fills:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(fills[(r - 1, c)])
                if fonts and (r - 1, c) in fonts:
                    color, bold = fonts[(r - 1, c)]
                    para.font.color.rgb = RGBColor.from_string(color)
                    para.font.bold = bold
        return table

    def percent(rate):
        return "-" if rate is None else f"{rate * 100:.2f}%"

    # ---- slide 1: title ----
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = f"Rapport Hebdo {label}"
    subtitle = (f"Tickets & Incidents - P1/P2 MTTR\n"
                f"Semaine du {metrics['week_start']} au {metrics['week_end']}\n"
                f"Sources : brise_smc_week{label[1:]} · GNOC_week{label[1:]}")
    title_slide.placeholders[1].text = subtitle

    # ---- slide 2: tickets ----
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = f"Tickets - {label}"
    headers = ["Domaine", "TT", "P1", "P2", "P3", "P4", "In SLA", "Hors SLA", "En cours"]
    data = []
    for name in metrics["row_order"]:
        m = rows[name]
        data.append([
            name,
            m["tt"] if m["tt"] is not None else "-",
            percent(m["rates"]["P1"]), percent(m["rates"]["P2"]),
            percent(m["rates"]["P3"]), percent(m["rates"]["P4"]),
            m["in_sla"] if m["in_sla"] is not None else "-",
            m["hors_sla"] if m["hors_sla"] is not None else "-",
            m["en_cours"] if m["en_cours"] is not None else "-",
        ])
    add_table(slide, headers, data, top=1.5,
              col_widths=[2.6, 0.9, 1.2, 1.2, 1.2, 1.2, 1.3, 1.4, 1.3])

    # ---- slides 3 & 4: MTTR ----
    for prio in MTTR_PRIORITIES:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"MTTR {prio} - {label}"
        note = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(12.0), Inches(0.35))
        note.text_frame.text = f"Cible SLA : {TARGET_HOURS[prio]:.0f} h"
        note.text_frame.paragraphs[0].font.size = Pt(13)
        note.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x66, 0x70, 0x85)
        headers = ["Domaine", "MTTR", "Ouvert"]
        data, fills, fonts = [], {}, {}
        for i, name in enumerate(metrics["row_order"]):
            m = rows[name]
            ouvert = m["ouvert"][prio]
            data.append([name, m["mttr"][prio] or "-",
                         ouvert if ouvert is not None else "-"])
            if m["mttr"][prio]:
                ok = m["mttr_ok"][prio]
                # Highlight: over-target cells (P1 > 4 h, P2 > 8 h) get a red
                # fill AND red bold text; within-target ones go green.
                fills[(i, 1)] = FILL_OK if ok else FILL_BAD
                fonts[(i, 1)] = (FONT_OK if ok else FONT_BAD, True)
        add_table(slide, headers, data, top=1.5,
                  col_widths=[4.0, 3.0, 3.0], fills=fills, fonts=fonts, font_size=13)

    buffer = io.BytesIO()
    prs.save(buffer)
    return f"hebdo_{label}.pptx", buffer.getvalue()


# =========================================================================== #
# One-shot generation
# =========================================================================== #

def generate(brise: Tuple[str, bytes], gnoc: Tuple[str, bytes],
             *, week: Any = None, date: Any = None,
             columns: Optional[Dict[str, str]] = None,
             outputs: Sequence[str] = ("xlsx", "pptx")
             ) -> Tuple[List[Tuple[str, bytes]], Dict[str, Any]]:
    """Analyse both sources and render the requested outputs."""
    metrics = analyze(brise, gnoc, week=week, date=date, columns=columns)
    files: List[Tuple[str, bytes]] = []
    if "xlsx" in outputs:
        files.append(render_excel(metrics))
    if "pptx" in outputs:
        files.append(render_pptx(metrics))
    if not files:
        raise HebdoError(f"Nothing to generate (outputs={list(outputs)!r}).")
    return files, metrics
