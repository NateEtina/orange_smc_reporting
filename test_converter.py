"""End-to-end tests for the CSV -> Excel engine.

Run:  python test_converter.py
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time

import converter
from converter import Options, convert_file, convert_many, preview_file

FAILURES: list[str] = []
CHECKS = 0


def check(condition, message: str) -> None:
    global CHECKS
    CHECKS += 1
    if not condition:
        FAILURES.append(message)


def load(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


# --------------------------------------------------------------------------
# Sample CSVs
# --------------------------------------------------------------------------

MAIN_CSV = (
    "id,name,qty,grouped,total,amount_eu,ratio,active,joined,dmy,last_seen,clock,"
    "sku,zip,big,formula,mixed\n"
    '1,"Café ""Premium""",10,"1,234","$1,234.56","1.234,56",45%,true,'
    "2024-01-15,15/01/2024,2024-01-15 09:30:00,09:30:00,007,01234,"
    '12345678901234567890,"=1+1",1\n'
    '2,"日本語 テキスト",20,"2,345","$99.00","99,99",12.5%,false,'
    "2024-02-20,20/02/2024,2024-02-20T10:00:00Z,13:45:00,008,05678,"
    '98765432109876543210,"=2+2",2\n'
    "3,Plain,30,\"10,000\",\"$2,500.00\",\"2.500,00\",3%,true,"
    "2024-03-05,05/03/2024,2024-03-05 23:05:00,23:05:00,009,99999,"
    '5,abc,x\n'
)

EU_CSV = "name;value;when\nAlice;1.234,56;01/02/2024\nBob;2.500,00;03/04/2024\n"

CP1252_CSV = "city;population\nKöln;1087000\nZürich;421000\n"


def run_tests() -> None:
    from openpyxl.utils import get_column_letter

    # ---------------- main file ------------------------------------------
    xlsx, report = convert_file("main.csv", MAIN_CSV.encode("utf-8"), Options())
    wb = load(xlsx)
    ws = wb[wb.sheetnames[0]]

    check(wb.sheetnames == ["main", "Details"], f"sheets: {wb.sheetnames}")
    check(ws.max_row == 4, f"rows {ws.max_row}")
    check(ws.max_column == 17, f"cols {ws.max_column}")

    # header styling / view options
    check(ws["A1"].font.bold is True, "header not bold")
    check(ws.freeze_panes == "A2", f"freeze {ws.freeze_panes}")
    check(ws.auto_filter.ref is not None, "autofilter missing")
    check(ws.column_dimensions["A"].width and ws.column_dimensions["A"].width >= 8,
          "column width missing")
    check(ws.row_dimensions[1].height == 20, "header row height")

    # typed values
    check(ws["A2"].value == 1 and isinstance(ws["A2"].value, int), f"id {ws['A2'].value!r}")
    check(ws["C2"].value == 10, f"qty {ws['C2'].value!r}")
    check(ws["D2"].value == 1234, f"grouped {ws['D2'].value!r}")
    check(ws["D2"].number_format == "#,##0", f"grouped fmt {ws['D2'].number_format}")
    check(ws["E2"].value == 1234.56, f"total {ws['E2'].value!r}")
    check("USD" not in ws["E2"].number_format and "$" in ws["E2"].number_format
          or '"$"' in ws["E2"].number_format, f"total fmt {ws['E2'].number_format}")
    check(ws["F2"].value == 1234.56, f"amount_eu {ws['F2'].value!r}")
    check(abs(ws["G2"].value - 0.45) < 1e-9, f"ratio {ws['G2'].value!r}")
    check(ws["G2"].number_format.endswith("%"), f"ratio fmt {ws['G2'].number_format}")
    check(ws["H2"].value is True, f"active {ws['H2'].value!r}")
    check(ws["I2"].value.date() == date(2024, 1, 15) if isinstance(ws["I2"].value, datetime)
          else ws["I2"].value == date(2024, 1, 15), f"joined {ws['I2'].value!r}")
    check(ws["J2"].value.date() == date(2024, 1, 15) if isinstance(ws["J2"].value, datetime)
          else ws["J2"].value == date(2024, 1, 15), f"dmy parsed {ws['J2'].value!r}")
    check(isinstance(ws["K2"].value, datetime) and ws["K2"].value.hour == 9,
          f"last_seen {ws['K2'].value!r}")
    check(ws["K2"].value.tzinfo is None, "timezone not stripped")
    check(ws["L2"].value in (time(9, 30), "09:30:00") or
          (isinstance(ws["L2"].value, time) and ws["L2"].value.hour == 9),
          f"clock {ws['L2'].value!r}")
    check(ws["M2"].value == "007", f"leading zeros sku {ws['M2'].value!r}")
    check(ws["N3"].value == "05678", f"leading zeros zip {ws['N3'].value!r}")
    check(ws["O2"].value == "12345678901234567890", f"big int {ws['O2'].value!r}")
    check(ws["O2"].data_type == "s", f"big int type {ws['O2'].data_type}")
    check(ws["P2"].value == "=1+1" and ws["P2"].data_type == "s",
          f"formula not escaped: {ws['P2'].value!r} {ws['P2'].data_type}")
    check(ws["Q4"].value == "x", f"mixed col {ws['Q4'].value!r}")
    check(ws["B2"].value == 'Café "Premium"', f"name {ws['B2'].value!r}")
    check(ws["B3"].value == "日本語 テキスト", f"cjk {ws['B3'].value!r}")

    # report
    check(report["rows"] == 3, f"report rows {report['rows']}")
    check(report["columns"] == 17, f"report cols {report['columns']}")
    types = {t["column"]: t["type"] for t in report["types"]}
    check(types["id"] == "integer", f"type id {types['id']}")
    check(types["total"] == "currency $", f"type total {types['total']}")
    check(types["ratio"] == "percent", f"type ratio {types['ratio']}")
    check(types["active"] == "boolean", f"type active {types['active']}")
    check(types["joined"] == "date", f"type joined {types['joined']}")
    check(types["last_seen"] == "datetime", f"type last_seen {types['last_seen']}")
    check(types["clock"] == "time", f"type clock {types['clock']}")
    check(types["sku"] == "text", f"type sku {types['sku']}")

    # details sheet
    details = wb["Details"]
    check(details["A1"].value == "CSV -> Excel conversion details", "details header")
    check(details.max_row >= 17, f"details rows {details.max_row}")

    # ---------------- semicolon + EU decimals + DMY/MDY ambiguity ----------
    xlsx_eu, report_eu = convert_file("eu.csv", EU_CSV.encode("utf-8"), Options())
    wb_eu = load(xlsx_eu)
    ws_eu = wb_eu["eu"]
    check(report_eu["delimiter"] == ";", f"delimiter {report_eu['delimiter']!r}")
    check(ws_eu["B2"].value == 1234.56, f"EU value {ws_eu['B2'].value!r}")
    # 01/02/2024 is ambiguous -> default MDY means Jan 2nd
    v = ws_eu["C2"].value
    d = v.date() if isinstance(v, datetime) else v
    check(d == date(2024, 1, 2), f"ambiguous date default {d}")
    # explicit DMY means Feb 1st
    xlsx_dmy, _ = convert_file("eu.csv", EU_CSV.encode("utf-8"), Options(date_order="DMY"))
    v2 = load(xlsx_dmy)["eu"]["C2"].value
    d2 = v2.date() if isinstance(v2, datetime) else v2
    check(d2 == date(2024, 2, 1), f"DMY date {d2}")

    # ---------------- encodings -------------------------------------------
    xlsx_cp, rep_cp = convert_file("cp.csv", CP1252_CSV.encode("cp1252"), Options())
    ws_cp = load(xlsx_cp)["cp"]
    check(rep_cp["encoding"] not in ("", None), "encoding not reported")
    check(ws_cp["A2"].value == "Köln", f"cp1252 value {ws_cp['A2'].value!r}")

    utf16 = "name;qty\nJosé;5\n".encode("utf-16")  # with BOM
    xlsx_u16, rep_u16 = convert_file("u16.csv", utf16, Options())
    check(load(xlsx_u16)["u16"]["A2"].value == "José", "utf-16 value")
    check("utf-16" in rep_u16["encoding"], f"utf-16 encoding {rep_u16['encoding']}")

    bom = "name,qty\nZoë,7\n".encode("utf-8-sig")
    _, rep_bom = convert_file("bom.csv", bom, Options())
    check("utf-8" in rep_bom["encoding"], f"bom encoding {rep_bom['encoding']}")

    # ---------------- forced delimiter / no header -------------------------
    no_header = "1,2,3\n4,5,6\n"
    _, rep_nh = convert_file("nh.csv", no_header.encode(), Options(has_header=False))
    wb_nh = load(convert_file("nh.csv", no_header.encode(), Options(has_header=False))[0])
    ws_nh = wb_nh["nh"]
    check(ws_nh["A1"].value == 1, f"no-header first cell {ws_nh['A1'].value!r}")
    check(ws_nh["A1"].font.bold is not True, "no-header sheet should have no styled header row")
    check(rep_nh["columns"] == 3, f"no-header cols {rep_nh['columns']}")

    # ---------------- ragged rows -----------------------------------------
    ragged = "a,b,c\n1,2\n3,4,5,6\n"
    xlsx_r, rep_r = convert_file("ragged.csv", ragged.encode(), Options())
    ws_r = load(xlsx_r)["ragged"]
    check(ws_r.max_column == 4, f"ragged width {ws_r.max_column}")
    check(any("padded" in w for w in rep_r["warnings"]), f"ragged warning {rep_r['warnings']}")

    # ---------------- empty rows & cells ----------------------------------
    sparse = "a,b\n1,\n,2\n\n3,4\n"
    xlsx_s, rep_s = convert_file("sparse.csv", sparse.encode(), Options())
    ws_s = load(xlsx_s)["sparse"]
    check(ws_s.max_row == 4, f"sparse rows {ws_s.max_row} (header + 3)")
    check(ws_s["B2"].value is None, f"sparse empty cell {ws_s['B2'].value!r}")
    check(any("empty row" in w for w in rep_s["warnings"]), f"sparse warning {rep_s['warnings']}")

    # ---------------- multi-file: one workbook ------------------------------
    outputs, reports = convert_many(
        [("main.csv", MAIN_CSV.encode()), ("eu.csv", EU_CSV.encode())],
        Options(),
    )
    check(len(outputs) == 1, f"single-workbook outputs {len(outputs)}")
    wb_multi = load(outputs[0][1])
    check(set(["main", "eu", "Details"]) <= set(wb_multi.sheetnames),
          f"multi sheets {wb_multi.sheetnames}")
    check(len(reports) == 2, f"multi reports {len(reports)}")

    # ---------------- multi-file: separate workbooks ------------------------
    outputs2, reports2 = convert_many(
        [("main.csv", MAIN_CSV.encode()), ("eu.csv", EU_CSV.encode())],
        Options(single_workbook=False, include_summary=False),
    )
    check(len(outputs2) == 2, f"separate outputs {len(outputs2)}")
    check([n for n, _ in outputs2] == ["main.xlsx", "eu.xlsx"], f"names {[n for n, _ in outputs2]}")
    check(len(reports2) == 2, f"separate reports {len(reports2)}")
    check("Details" not in load(outputs2[0][1]).sheetnames, "summary should be off")

    # ---------------- sheet splitting at Excel row limit --------------------
    original = converter.MAX_ROWS_TOTAL
    converter.MAX_ROWS_TOTAL = 4  # header + 3 data rows per sheet
    try:
        many_rows = "id,v\n" + "".join(f"{i},{i * 2}\n" for i in range(1, 11))
        xlsx_split, rep_split = convert_file("split.csv", many_rows.encode(), Options())
    finally:
        converter.MAX_ROWS_TOTAL = original
    wb_split = load(xlsx_split)
    data_sheets = [s for s in wb_split.sheetnames if s != "Details"]
    check(len(data_sheets) == 4, f"split sheets {data_sheets}")
    check(any("split across" in w for w in rep_split["warnings"]), f"split warning {rep_split['warnings']}")
    total = sum(wb_split[s].max_row - 1 for s in data_sheets)
    check(total == 10, f"split total rows {total}")
    check(wb_split[data_sheets[1]]["A1"].value == "id", "continuation sheet header")

    # ---------------- styling options ---------------------------------------
    opts = Options(
        style_header=True, header_fill="FF0000", header_font="00FF00",
        borders=True, zebra=True, wrap_text=True, freeze_panes=False,
        auto_filter=False, auto_width=False,
    )
    xlsx_styled, _ = convert_file("styled.csv", MAIN_CSV.encode(), opts)
    ws_st = load(xlsx_styled)["styled"]
    check(ws_st["A1"].fill.fgColor.rgb is not None, "header fill missing")
    check("FF0000" in str(ws_st["A1"].fill.fgColor.rgb), f"fill {ws_st['A1'].fill.fgColor.rgb}")
    check("00FF00" in str(ws_st["A1"].font.color.rgb), f"font color {ws_st['A1'].font.color.rgb}")
    check(ws_st.freeze_panes is None, "freeze should be off")
    check(ws_st.auto_filter.ref is None, "autofilter should be off")
    check(ws_st["A2"].border.left.style == "thin", "borders missing")
    check(ws_st["A3"].fill.fgColor.rgb is not None, "zebra missing")
    check(ws_st["B2"].alignment.wrap_text is True, "wrap missing")

    # ---------------- escaping off -------------------------------------------
    xlsx_raw, _ = convert_file("raw.csv", MAIN_CSV.encode(), Options(escape_formulas=False))
    cell = load(xlsx_raw)["raw"]["P2"]
    check(cell.data_type == "f", f"formula should stay a formula, got {cell.data_type}")

    # ---------------- preview -------------------------------------------------
    prev = preview_file("main.csv", MAIN_CSV.encode(), Options())
    check(prev["rows"] == 3, f"preview rows {prev['rows']}")
    check(len(prev["header"]) == 17, f"preview cols {len(prev['header'])}")
    check(len(prev["preview"]) == 3, f"preview len {len(prev['preview'])}")
    ptypes = {t["column"]: t["type"] for t in prev["types"]}
    check(ptypes["ratio"] == "percent", f"preview type {ptypes['ratio']}")

    # ---------------- formula-heavy & big numbers -----------------------------
    edge = "v,w\n=SUM(A1:A9),1234567890123456789012345\n@cmd,20\n+1234,30\nx,40\n"
    xlsx_e, rep_e = convert_file("edge.csv", edge.encode(), Options())
    ws_e = load(xlsx_e)["edge"]
    check(ws_e["A2"].data_type == "s" and ws_e["A2"].value == "=SUM(A1:A9)",
          f"edge formula {ws_e['A2'].value!r} {ws_e['A2'].data_type}")
    check(ws_e["B2"].value == "1234567890123456789012345", f"edge big {ws_e['B2'].value!r}")
    check(ws_e["B2"].data_type == "s", f"edge big type {ws_e['B2'].data_type}")
    check(any("kept as text" in w for w in rep_e["warnings"]), f"edge warning {rep_e['warnings']}")

    # ---------------- zip packaging (as done by the API) ----------------------
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in outputs2:
            zf.writestr(name, data)
    check(zipfile.ZipFile(io.BytesIO(buf.getvalue())).namelist() == ["main.xlsx", "eu.xlsx"],
          "zip contents")

    # ---------------- monitoring profile: key=value tags ----------------------
    check(converter.parse_tags("a=1, b=2 ,c=x\\,y,b=3") == {"a": "1", "b": "3", "c": "x,y"},
          f"parse_tags {converter.parse_tags('a=1, b=2 ,c=x\\,y,b=3')}")
    check(converter.parse_tags("") == {} and converter.parse_tags("noequals") == {},
          "parse_tags empty/no-key")
    check(converter._unescape("Alger\\ Centre") == "Alger Centre", "unescape backslash-space")
    check(converter._epoch_to_datetime("1577836800000000000") == datetime(2020, 1, 1),
          f"epoch ns {converter._epoch_to_datetime('1577836800000000000')}")
    check(converter._epoch_to_datetime("1577836800000") == datetime(2020, 1, 1),
          "epoch ms auto-detect")
    check(converter._epoch_to_datetime("abc") is None and converter._epoch_to_datetime("5") is None,
          "epoch rejects non-timestamps")

    # renaming: tag keys + English synonyms -> the French report headers
    check(converter.rename_monitoring_columns(
        ["time", "name", "count_duration_seconds", "tags", "zone", "Status",
         "monitoring_name", "cellular_sms_reception_status", "unknown_col"])
        == ["Horodatage (UTC)", "Type Enregistrement", "Durée (sec)", "tags",
            "Zone Géographique", "Statut Global", "Nom du Service (Monitoring)",
            "Statut SMS", "unknown_col"],
        "rename to French headers")
    check(converter.rename_monitoring_columns(["status", "Status"])
          == ["Statut Global", "Statut Global (2)"], "rename dedups collisions")

    rh, rd = converter.reorder_monitoring_columns(
        ["extra", "Nom du Test", "Statut Global", "Date"],
        [["x", "ping", "success", "2024-01-02"]])
    check(rh == ["Date", "Statut Global", "Nom du Test", "extra"], f"reorder {rh}")
    check(rd == [["2024-01-02", "success", "ping", "x"]], f"reorder rows {rd}")

    # ---------------- monitoring profile: end-to-end --------------------------
    MONITOR_CSV = (
        "name,tags,time,count_duration_seconds\n"
        'test,"zone=Alger\\ Centre,status=success,failure_reason=,monitoring_name=Svc A",1577836800000000000,1\n'
        'test,"zone=Oran,status=failed,failure_reason=Timeout,monitoring_name=Svc B,'
        'cellular_sms_reception_status=failed",1580515200000000000,2\n'
        'test,"status=error",,3\n'
    )
    opts_m = Options(monitoring_profile=True, include_summary=False)
    check(opts_m.zebra is True and opts_m.zebra_fill == "F9FAFB",
          f"profile forces zebra {opts_m.zebra} {opts_m.zebra_fill}")
    xlsx_m, rep_m = convert_file("monitor.csv", MONITOR_CSV.encode("utf-8"), opts_m)
    wb_m = load(xlsx_m)
    check(wb_m.sheetnames == ["Detailed Raw Data"], f"monitor sheet {wb_m.sheetnames}")
    ws_m = wb_m["Detailed Raw Data"]

    check([ws_m.cell(1, c).value for c in range(1, 10)] ==
          ["Date", "Horodatage (UTC)", "Zone Géographique", "Statut Global",
           "Nom du Service (Monitoring)", "Raison de l'Échec", "Durée (sec)",
           "Statut SMS", "Type Enregistrement"],
          f"monitor header {[ws_m.cell(1, c).value for c in range(1, 10)]}")
    check("FF1F4E78" == str(ws_m["A1"].fill.fgColor.rgb), f"header fill {ws_m['A1'].fill.fgColor.rgb}")
    check(ws_m["A1"].font.bold is True and str(ws_m["A1"].font.color.rgb) == "FFFFFFFF",
          f"header font {ws_m['A1'].font.color.rgb}")
    check(ws_m["A1"].font.name == "Segoe UI" and ws_m["A1"].font.sz == 11, "header font family")
    check(ws_m.freeze_panes == "A2" and ws_m.auto_filter.ref is not None,
          "profile freeze/autofilter")

    # row 1: success -> green, escaped tag value, empty failure -> label
    check((ws_m["A2"].value.date() if isinstance(ws_m["A2"].value, datetime)
           else ws_m["A2"].value) == date(2020, 1, 1), f"Date cell {ws_m['A2'].value!r}")
    check(ws_m["B2"].value == datetime(2020, 1, 1), f"Timestamp cell {ws_m['B2'].value!r}")
    check(ws_m["C2"].value == "Alger Centre", f"tag unescape {ws_m['C2'].value!r}")
    check(ws_m["D2"].value == "success" and str(ws_m["D2"].fill.fgColor.rgb) == "FFE2EFDA",
          f"success cell {ws_m['D2'].value!r} {ws_m['D2'].fill.fgColor.rgb}")
    check(ws_m["E2"].value == "Svc A", f"service name {ws_m['E2'].value!r}")
    check(ws_m["F2"].value == "Aucune (Succès)", f"failure label {ws_m['F2'].value!r}")
    check(ws_m["G2"].value == 1, f"duration {ws_m['G2'].value!r}")
    check(ws_m["H2"].value == "N/A", f"sms label {ws_m['H2'].value!r}")

    # row 2: failed -> red, zebra stripe on the other cells (F9FAFB)
    check(ws_m["D3"].value == "failed" and str(ws_m["D3"].fill.fgColor.rgb) == "FFFCE4D6",
          f"failed cell {ws_m['D3'].value!r} {ws_m['D3'].fill.fgColor.rgb}")
    check(str(ws_m["A3"].fill.fgColor.rgb) == "FFF9FAFB", f"zebra {ws_m['A3'].fill.fgColor.rgb}")
    check(ws_m["F3"].value == "Timeout", f"failure reason {ws_m['F3'].value!r}")

    # row 3: no time -> empty Date/Timestamp, empty zone -> 'Non Définie', error -> yellow
    check(ws_m["A4"].value is None and ws_m["B4"].value is None, "empty timestamp kept empty")
    check(ws_m["C4"].value == "Non Définie", f"zone label {ws_m['C4'].value!r}")
    check(ws_m["D4"].value == "error" and str(ws_m["D4"].fill.fgColor.rgb) == "FFFFF2CC",
          f"error cell {ws_m['D4'].value!r} {ws_m['D4'].fill.fgColor.rgb}")
    check(ws_m["D2"].font.name == "Segoe UI" and ws_m["D2"].font.sz == 9
          and ws_m["D2"].font.bold is True, f"status font {ws_m['D2'].font.name}")
    check(ws_m["C2"].font.name == "Segoe UI" and ws_m["C2"].font.sz == 9
          and not ws_m["C2"].font.bold, f"body font {ws_m['C2'].font.name}")

    check(any("expanded into 5 key=value" in w for w in rep_m["warnings"]),
          f"tags warning {rep_m['warnings']}")
    check(any("UNIX timestamp" in w for w in rep_m["warnings"]),
          f"time warning {rep_m['warnings']}")

    # double-quoted ("wrapped") lines are unwrapped back into columns
    wrapped = (
        'name,tags,time,count_duration_seconds\n'
        '"test,""zone=Kinshasa,status=success,failure_reason=Robot\\ unreachable"",'
        '1577836800000000000,0"\n'
        '"test,""zone=Goma,status=failed,failure_reason=Timeout"",1580515200000000000,1"\n'
    )
    xlsx_wr, _ = convert_file("wrapped.csv", wrapped.encode("utf-8"), opts_m)
    ws_wr = load(xlsx_wr)["Detailed Raw Data"]
    check(ws_wr.max_column == 7, f"wrapped columns {ws_wr.max_column}")
    check(ws_wr["C2"].value == "Kinshasa" and ws_wr["C3"].value == "Goma",
          f"wrapped zone {ws_wr['C2'].value!r} {ws_wr['C3'].value!r}")
    check(ws_wr["E2"].value == "Robot unreachable", f"wrapped unescape {ws_wr['E2'].value!r}")
    check(ws_wr["B3"].value == datetime(2020, 2, 1), f"wrapped time {ws_wr['B3'].value!r}")

    # a normal file must NOT be unwrapped
    plain_rows = converter.read_rows("a,b\n1,2\n3,4\n", ",", '"')
    check(plain_rows == [["a", "b"], ["1", "2"], ["3", "4"]], f"read_rows normal {plain_rows}")

    # unparsed 'time' values stay untouched and are reported
    weird = 'time,tags\nabc,"status=success"\n1577836800000000000,\n'
    xlsx_w, rep_w = convert_file("weird.csv", weird.encode(), opts_m)
    ws_w = load(xlsx_w)["Detailed Raw Data"]
    check(ws_w["B2"].value == "abc", f"unparsed time {ws_w['B2'].value!r}")
    check(any("not UNIX timestamps" in w for w in rep_w["warnings"]),
          f"unparsed warning {rep_w['warnings']}")

    # profile with no tags/time column is a no-op apart from the sheet name
    xlsx_n, rep_n = convert_file("plain.csv", b"k,v\n1,2\n3,4\n", opts_m)
    check(load(xlsx_n).sheetnames == ["Detailed Raw Data"], "profile sheet without tags/time")
    check(not any("tags" in w for w in rep_n["warnings"]), f"no-tags warnings {rep_n['warnings']}")

    # preview reflects the transformed layout
    prev_m = preview_file("monitor.csv", MONITOR_CSV.encode("utf-8"), Options(monitoring_profile=True))
    check(prev_m["header"][:2] == ["Date", "Horodatage (UTC)"], f"preview header {prev_m['header']}")
    check("Zone Géographique" in prev_m["header"], "preview tag columns")

    # from_dict forces the mandated look even against explicit switches
    forced = Options.from_dict({"monitoring_profile": True, "zebra": False,
                                "header_fill": "FF0000", "freeze_panes": False,
                                "zebra_fill": "ABCDEF"})
    check(forced.zebra is True and forced.zebra_fill == "F9FAFB", "from_dict forces zebra")
    check(forced.header_fill == "1F4E78" and forced.freeze_panes is True,
          f"from_dict forces header/freeze {forced.header_fill} {forced.freeze_panes}")

    # custom zebra color outside the profile
    xlsx_z, _ = convert_file("plain.csv", b"k,v\n1,2\n3,4\n5,6\n",
                             Options(zebra=True, zebra_fill="ABCDEF", include_summary=False))
    check("ABCDEF" in str(load(xlsx_z)["plain"]["A3"].fill.fgColor.rgb),
          "custom zebra color")

    # ---------------- monitoring profile: many files --------------------------
    out_one, _ = convert_many(
        [("a.csv", MONITOR_CSV.encode()), ("b.csv", MONITOR_CSV.encode())],
        Options(monitoring_profile=True, include_summary=False),
    )
    check(len(out_one) == 1, f"monitor single-workbook outputs {len(out_one)}")
    check(load(out_one[0][1]).sheetnames == ["Detailed Raw Data", "Detailed Raw Data (2)"],
          f"monitor multi sheets {load(out_one[0][1]).sheetnames}")

    out_many, _ = convert_many(
        [("a.csv", MONITOR_CSV.encode()), ("b.csv", MONITOR_CSV.encode())],
        Options(monitoring_profile=True, single_workbook=False, include_summary=False),
    )
    check(len(out_many) == 2 and [n for n, _ in out_many] == ["a.xlsx", "b.xlsx"],
          f"monitor separate outputs {[n for n, _ in out_many]}")
    check(load(out_many[0][1]).sheetnames == ["Detailed Raw Data"],
          "each separate workbook uses the profile sheet name")


if __name__ == "__main__":
    run_tests()
    if FAILURES:
        print(f"FAILED {len(FAILURES)}/{CHECKS} checks:")
        for failure in FAILURES:
            print(f"  - {failure}")
        raise SystemExit(1)
    print(f"OK: all {CHECKS} checks passed")
